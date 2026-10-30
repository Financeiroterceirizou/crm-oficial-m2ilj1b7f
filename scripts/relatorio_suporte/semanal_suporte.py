#!/usr/bin/env python3
# -*- coding: utf-8 -*-
# RELATÓRIO SEMANAL — SUPORTE PROTEÇÃO VEICULAR — v1.0, 2026-10-30 (pedido Vinícius)
#   Toda sexta-feira:
#   (1) Consolidado PAGO do dia 01 até a Nª sexta-feira do mês corrente (por categoria);
#   (2) GERAL de tudo que tem A PAGAR (não consolidado) até o fim do mês corrente;
#   (3) Contas a pagar do MÊS SEGUINTE (lançamento a lançamento).
#   Regime de caixa; transferências entre contas fora.
# Uso: python3 semanal_rede_sos.py [YYYY-MM-DD] [--enviar]
import json, os, sys, urllib.request, time, base64, hashlib
from collections import defaultdict
from datetime import date, timedelta
from reportlab.lib.pagesizes import A4
from reportlab.lib.units import cm
from reportlab.lib import colors
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, PageBreak
from reportlab.platypus import Image as RLImage
from PIL import Image as PILImage
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

BASE = "https://api-v1.controlle.com"
_dir = os.path.dirname(os.path.abspath(__file__))
TOKEN = os.environ.get("CONTROLLE_TOKEN_SUPORTE") or open(os.path.join(_dir, ".controlle_token")).read().strip()
UA = {"Authorization": f"Bearer {TOKEN}", "User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) Chrome/126.0"}
LARANJA = colors.HexColor("#ff501c"); LARANJA_CLARO = colors.HexColor("#ffe3d6")
PRETO = colors.HexColor("#1a1a1a"); FUNDO_HDR = colors.HexColor("#e8e8e8")
VERDE = colors.HexColor("#1a7f37"); VERMELHO = colors.HexColor("#c0392b")
LOGO_SUP = os.path.join(_dir, "logo-suporte.png")
LOGO_TZ = os.path.join(_dir, "logo-terceirizou.png")
SIMBOLO = os.path.join(_dir, "simbolo-terceirizou.png")
MES_PT = {1: "janeiro", 2: "fevereiro", 3: "março", 4: "abril", 5: "maio", 6: "junho",
          7: "julho", 8: "agosto", 9: "setembro", 10: "outubro", 11: "novembro", 12: "dezembro"}
TRANSF_CONTAS = {10220988, 10491428, 10221133, 10221134}

def req(url, tries=4):
    for i in range(tries):
        try:
            r = urllib.request.Request(url)
            for k, v in UA.items():
                r.add_header(k, v)
            with urllib.request.urlopen(r, timeout=120) as resp:
                return json.loads(resp.read().decode("utf-8", "replace"))
        except urllib.error.HTTPError as e:
            if e.code in (500, 502, 503, 409) and i < tries - 1:
                time.sleep(4 * (i + 1)); continue
            raise

def tx_list(start, end, **filtros):
    out, page = [], 1
    while True:
        url = (f"{BASE}/transaction/v1/transactions/list?start_date={start}&end_date={end}"
               f"&page={page}&orderBy=date&orderByCardinality=ASC")
        for k, v in filtros.items():
            url += f"&{k}={v}"
        tl = req(url).get("results", {}).get("transactionsList", [])
        out.extend(tl)
        if len(tl) < 100:
            return out
        page += 1

def bdate(t):
    return (t.get("dt_billing") or t.get("dt_due") or "")[:10]

def carregar_plano():
    plano = [p for p in req(f"{BASE}/plan-account/v1/planAccountsEntities").get("results", [])
             if p.get("status") == 1]
    by_id = {p["id"]: p for p in plano}
    raiz_de = {}
    for p in plano:
        cid, vistos = p["id"], set()
        while cid in by_id and by_id[cid].get("id_plan_accounts_parent") and cid not in vistos:
            vistos.add(cid)
            cid = by_id[cid]["id_plan_accounts_parent"]
        raiz_de[p["id"]] = cid
    return by_id, raiz_de

BY_ID, RAIZ_DE = carregar_plano()

def raiz(cid):
    return RAIZ_DE.get(cid, cid)

def cat_raiz_nome(t):
    """Nome da categoria-raiz do lançamento (para agrupar)."""
    ids = {raiz(c.get("id_category")) for c in (t.get("apportionments_plan_account") or []) if c.get("id_category")}
    ids -= TRANSF_CONTAS
    if not ids:
        return "(sem categoria)"
    return " + ".join(sorted((BY_ID.get(i, {}).get("ds_category") or str(i)) for i in ids))

def eh_transf(t):
    ids = {raiz(c.get("id_category")) for c in (t.get("apportionments_plan_account") or []) if c.get("id_category")}
    return bool(ids) and ids <= TRANSF_CONTAS

def valor_efetivo(t):
    if t.get("situation") in (1, 2) and t.get("payment_in_cent") is not None:
        return t["payment_in_cent"]
    return t["value_in_cent"]

def brl(cents):
    v = cents / 100
    s = f"{abs(v):,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")
    return ("-" if v < 0 else "") + "R$ " + s

# ===== parâmetros =====
HOJE = date.fromisoformat(sys.argv[1]) if len(sys.argv) > 1 and not sys.argv[1].startswith("--") else date.today()
ENVIAR = "--enviar" in sys.argv
MES_REF = HOJE.replace(day=1)
FIM_MES_REF = (MES_REF.replace(day=28) + timedelta(days=4)).replace(day=1) - timedelta(days=1)
MES_SEG = (FIM_MES_REF + timedelta(days=1)).replace(day=1)
FIM_MES_SEG = (MES_SEG.replace(day=28) + timedelta(days=4)).replace(day=1) - timedelta(days=1)
# nª sexta-feira do mês (qual sexta é hoje)
_sextas = [d for d in range(1, (FIM_MES_REF.day + 1)) if date(MES_REF.year, MES_REF.month, d).weekday() == 4]
_passadas = [d for d in _sextas if d <= HOJE.day]
n_sexta = len(_passadas) or 1
DIA_REF = max(_passadas) if _passadas else HOJE.day

# ===== dados =====
# (1) consolidado pago: dia 01 → Nª sexta
tx_cons = [t for t in tx_list(MES_REF.isoformat(), date(MES_REF.year, MES_REF.month, DIA_REF).isoformat(),
                              **{"situation": "[1,2]"})
           if not eh_transf(t)]
por_cat_cons = defaultdict(lambda: [0, 0])
for t in tx_cons:
    v = valor_efetivo(t)
    c = cat_raiz_nome(t)
    por_cat_cons[c][0 if v > 0 else 1] += v

# (2) a pagar até o fim do mês corrente (não consolidado)
tx_apagar_mes = [t for t in tx_list(HOJE.isoformat(), FIM_MES_REF.isoformat(), **{"situation": "[0]"})
                 if not eh_transf(t) and t.get("value_in_cent", 0) < 0]
# (3) contas a pagar do mês seguinte
tx_apagar_seg = [t for t in tx_list(MES_SEG.isoformat(), FIM_MES_SEG.isoformat(), **{"situation": "[0]"})
                 if not eh_transf(t) and t.get("value_in_cent", 0) < 0]

# ===== estilos =====
styles = getSampleStyleSheet()
h1c = ParagraphStyle("h1c", parent=styles["Heading1"], fontSize=13.5, alignment=1, textColor=PRETO, spaceAfter=2)
subc = ParagraphStyle("subc", parent=styles["Normal"], fontSize=7.2, alignment=1, textColor=colors.HexColor("#777777"), spaceAfter=4)
h2 = ParagraphStyle("h2", parent=styles["Heading2"], fontSize=10, textColor=LARANJA, spaceBefore=8, spaceAfter=3)
h3s = ParagraphStyle("h3s", parent=styles["Normal"], fontName="Helvetica-Bold", fontSize=8.4, spaceBefore=4, spaceAfter=1.5)
cell = ParagraphStyle("cell", parent=styles["Normal"], fontSize=7.4, leading=8.8)
cellb = ParagraphStyle("cellb", parent=cell, fontName="Helvetica-Bold")
cellr = ParagraphStyle("cellr", parent=cell, alignment=2)
cellrb = ParagraphStyle("cellrb", parent=cellb, alignment=2)
cellc = ParagraphStyle("cellc", parent=cell, alignment=1)
cellh = ParagraphStyle("cellh", parent=cell, fontName="Helvetica-Bold", fontSize=7.4)
cellch = ParagraphStyle("cellch", parent=cellh, alignment=1)
cellwr = ParagraphStyle("cellwr", parent=cellh, alignment=2)
P = Paragraph

def P_val(cents, style, bold=False):
    s = ParagraphStyle("v", parent=style, textColor=(VERDE if cents > 0 else VERMELHO if cents < 0 else PRETO))
    if bold: s.fontName = "Helvetica-Bold"
    return P(brl(cents), s)

def tabela(rows, widths, hdr_fundo=FUNDO_HDR, hdr_cor=PRETO, fs=7.4):
    t = Table(rows, colWidths=widths, repeatRows=1)
    st = [("FONTNAME", (0, 0), (-1, -1), "Helvetica"), ("FONTSIZE", (0, 0), (-1, -1), fs),
          ("TOPPADDING", (0, 0), (-1, -1), 2.2), ("BOTTOMPADDING", (0, 0), (-1, -1), 2.2),
          ("LEFTPADDING", (0, 0), (-1, -1), 3), ("RIGHTPADDING", (0, 0), (-1, -1), 3),
          ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
          ("BACKGROUND", (0, 0), (-1, 0), hdr_fundo),
          ("GRID", (0, 0), (-1, -1), 0.35, colors.HexColor("#d9d9d9")),
          ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#fafafa")])]
    t.setStyle(TableStyle(st))
    return t

def rodape(canvas, doc_):
    canvas.saveState()
    if os.path.exists(SIMBOLO):
        try:
            w0, h0 = PILImage.open(SIMBOLO).size
            canvas.drawImage(SIMBOLO, 1.3 * cm, 0.8 * cm, width=0.5 * cm, height=0.5 * cm * h0 / w0, mask="auto")
        except Exception:
            pass
    canvas.setFont("Helvetica", 6.5); canvas.setFillColor(colors.HexColor("#777777"))
    canvas.drawRightString(A4[0] - 1.4 * cm, 0.9 * cm, f"Página {doc_.page}")
    canvas.restoreState()

ARQ_PDF = f"artifacts/{HOJE.strftime('%y%m%d')}_Semanal_Suporte_{HOJE.strftime('%Y%m%d')}.pdf"
doc = SimpleDocTemplate(ARQ_PDF, pagesize=A4, leftMargin=1.5 * cm, rightMargin=1.5 * cm,
                        topMargin=1.2 * cm, bottomMargin=1.6 * cm,
                        title=f"Relatório Semanal — Suporte Proteção Veicular — {HOJE.strftime('%d/%m/%Y')}")
E = []
# cabeçalho: 2 logos + título
logos = []
if os.path.exists(LOGO_SUP):
    w0, h0 = PILImage.open(LOGO_SUP).size
    logos.append(RLImage(LOGO_SUP, width=4.2 * cm, height=4.2 * cm * h0 / w0))
if os.path.exists(LOGO_TZ):
    wt, ht = PILImage.open(LOGO_TZ).size
    logos.append(RLImage(LOGO_TZ, width=3.8 * cm, height=3.8 * cm * ht / wt))
if logos:
    cab = Table([[l for l in logos[0:1]] + [""] + [[l for l in logos[1:2]]]],
                colWidths=[3 * cm, 11 * cm, 3.9 * cm]) if len(logos) == 2 else None
    if cab:
        cab.setStyle(TableStyle([("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                                 ("ALIGN", (2, 0), (2, 0), "RIGHT")]))
        E.append(cab)
E.append(Spacer(1, 4))
E.append(P("<b>SUPORTE PROTEÇÃO VEICULAR — RELATÓRIO SEMANAL</b>", h1c))
E.append(P(f"{HOJE.strftime('%d/%m/%Y')} · {n_sexta}ª sexta-feira de {MES_PT[MES_REF.month]}/{MES_REF.year} "
           f"· Fonte: Controlle · Terceirizou", subc))

# ===== 1) consolidado pago =====
E.append(P(f"Consolidado Pago — 01/{MES_REF.month:02d} a {DIA_REF:02d}/{MES_REF.month:02d} ({n_sexta}ª sexta)", h2))
rows = [[P("<b>Categoria</b>", cellh), P("<b>Entradas</b>", cellwr), P("<b>Saídas</b>", cellwr), P("<b>Resultado</b>", cellwr)]]
tot_e = tot_s = 0
for cat in sorted(por_cat_cons):
    e_c, s_c = por_cat_cons[cat]
    tot_e += e_c; tot_s += s_c
    rows.append([P(cat, cell), P_val(e_c, cellr), P_val(s_c, cellr), P_val(e_c + s_c, cellrb)])
rows.append([P("<b>Total</b>", cellrb), P_val(tot_e, cellrb, True), P_val(tot_s, cellrb, True), P_val(tot_e + tot_s, cellrb, True)])
t = tabela(rows, [7.6 * cm, 3.4 * cm, 3.4 * cm, 3.4 * cm])
t.setStyle(TableStyle([("BACKGROUND", (0, len(rows) - 1), (-1, len(rows) - 1), LARANJA_CLARO)]))
E.append(t)

# ===== 2) a pagar até o fim do mês corrente =====
tot_ap_mes = sum(valor_efetivo(t) for t in tx_apagar_mes)
E.append(P(f"A Pagar — de hoje até {FIM_MES_REF.strftime('%d/%m/%Y')} (não consolidado)", h2))
por_cat_d = defaultdict(list)
for t in tx_apagar_mes:
    por_cat_d[cat_raiz_nome(t)].append(t)
for cat in sorted(por_cat_d):
    txs_c = sorted(por_cat_d[cat], key=bdate)
    rows = [[P("<b>Vencimento</b>", cellh), P("<b>Descrição</b>", cellh), P("<b>Conta</b>", cellh), P("<b>Valor</b>", cellwr)]]
    for t in txs_c:
        d = bdate(t)
        rows.append([P(f"{d[8:10]}/{d[5:7]}/{d[:4]}", cell), P((t.get("ds_transaction") or "")[:62], cell),
                     P((t.get("ds_account_main") or "")[:22], cell), P_val(valor_efetivo(t), cellr)])
    rows.append([P("<b>Total</b>", cellrb), P(f"<b>{cat}</b>", cellrb), P(f"<b>{len(txs_c)}</b>", cellc),
                 P_val(sum(valor_efetivo(t) for t in txs_c), cellrb)])
    tt = tabela(rows, [2.1 * cm, 8.0 * cm, 3.4 * cm, 3.3 * cm])
    tt.setStyle(TableStyle([("BACKGROUND", (0, len(rows) - 1), (-1, len(rows) - 1), LARANJA_CLARO)]))
    E.append(P(f"<b>{cat}</b>", h3s))
    E.append(tt)
rows = [[P("<b>Total a pagar no mês corrente</b>", cellrb), P(f"<b>{len(tx_apagar_mes)} lançamentos</b>", cellc),
         P_val(tot_ap_mes, cellrb, True)]]
t = tabela(rows, [9.8 * cm, 4 * cm, 4 * cm])
t.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, 0), LARANJA_CLARO)]))
E.append(Spacer(1, 3))
E.append(t)

# ===== 3) contas a pagar do mês seguinte =====
E.append(PageBreak())
tot_ap_seg = sum(valor_efetivo(t) for t in tx_apagar_seg)
E.append(P(f"Contas a Pagar — {MES_PT[MES_SEG.month].capitalize()} de {MES_SEG.year} (não consolidado)", h2))
por_cat_s = defaultdict(list)
for t in tx_apagar_seg:
    por_cat_s[cat_raiz_nome(t)].append(t)
if not tx_apagar_seg:
    E.append(P("Nenhum lançamento a pagar no mês seguinte.", cell))
for cat in sorted(por_cat_s):
    txs_c = sorted(por_cat_s[cat], key=bdate)
    rows = [[P("<b>Vencimento</b>", cellh), P("<b>Descrição</b>", cellh), P("<b>Conta</b>", cellh), P("<b>Valor</b>", cellwr)]]
    for t in txs_c:
        d = bdate(t)
        rows.append([P(f"{d[8:10]}/{d[5:7]}/{d[:4]}", cell), P((t.get("ds_transaction") or "")[:62], cell),
                     P((t.get("ds_account_main") or "")[:22], cell), P_val(valor_efetivo(t), cellr)])
    rows.append([P("<b>Total</b>", cellrb), P(f"<b>{cat}</b>", cellrb), P(f"<b>{len(txs_c)}</b>", cellc),
                 P_val(sum(valor_efetivo(t) for t in txs_c), cellrb)])
    tt = tabela(rows, [2.1 * cm, 8.0 * cm, 3.4 * cm, 3.3 * cm])
    tt.setStyle(TableStyle([("BACKGROUND", (0, len(rows) - 1), (-1, len(rows) - 1), LARANJA_CLARO)]))
    E.append(P(f"<b>{cat}</b>", h3s))
    E.append(tt)
rows = [[P(f"<b>Total a pagar — {MES_PT[MES_SEG.month].capitalize()}/{MES_SEG.year}</b>", cellrb),
         P(f"<b>{len(tx_apagar_seg)} lançamentos</b>", cellc), P_val(tot_ap_seg, cellrb, True)]]
t = tabela(rows, [9.8 * cm, 4 * cm, 4 * cm])
t.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, 0), LARANJA_CLARO)]))
E.append(Spacer(1, 3))
E.append(t)
E.append(Spacer(1, 8))
E.append(P("Gerado automaticamente pela Terceirizou · dados do Controlle", subc))
doc.build(E, onFirstPage=rodape, onLaterPages=rodape)
print(f"OK: {ARQ_PDF}")

# ===== Excel =====
ARQ_XLSX = f"artifacts/{HOJE.strftime('%y%m%d')}_Semanal_Suporte_{HOJE.strftime('%Y%m%d')}.xlsx"
wb = Workbook()
FILL_T = PatternFill("solid", fgColor="FFE3D6"); FILL_H = PatternFill("solid", fgColor="E8E8E8")
FH = Font(bold=True)

def cab(ws, titulo, cols):
    ws.append((titulo,))
    ws.cell(row=ws.max_row, column=1).font = Font(bold=True, size=11, color="FF501C")
    ws.append(cols)
    r = ws.max_row
    for c in range(1, len(cols) + 1):
        cell_ = ws.cell(row=r, column=c)
        cell_.font = FH; cell_.fill = FILL_H
        cell_.alignment = Alignment(horizontal="center" if c > 1 else "left")

ws = wb.active; ws.title = "Consolidado"
ws.append((f"SUPORTE PROTEÇÃO VEICULAR — Semanal — {n_sexta}ª sexta ({HOJE.strftime('%d/%m/%Y')}) — regime de caixa",))
ws["A1"].font = Font(bold=True, size=12)
cab(ws, f"Consolidado Pago — 01 a {DIA_REF:02d}/{MES_REF.month:02d}", ("Categoria", "Entradas", "Saídas", "Resultado"))
for cat in sorted(por_cat_cons):
    e_c, s_c = por_cat_cons[cat]
    ws.append((cat, e_c / 100, s_c / 100, (e_c + s_c) / 100))
ws.append(("Total", tot_e / 100, tot_s / 100, (tot_e + tot_s) / 100))
for c in range(1, 5):
    ws.cell(row=ws.max_row, column=c).font = FH
    ws.cell(row=ws.max_row, column=c).fill = FILL_T
ws.column_dimensions["A"].width = 38
for c in range(2, 5):
    ws.column_dimensions[get_column_letter(c)].width = 15

ws2 = wb.create_sheet("A Pagar (mês)")
cab(ws2, f"A Pagar — hoje até {FIM_MES_REF.strftime('%d/%m/%Y')}", ("Vencimento", "Descrição", "Conta", "Categoria", "Valor"))
for t in sorted(tx_apagar_mes, key=bdate):
    ws2.append((bdate(t), t.get("ds_transaction"), t.get("ds_account_main"), cat_raiz_nome(t), valor_efetivo(t) / 100))
ws2.append(("Total", "", "", "", tot_ap_mes / 100))
ws2.cell(row=ws2.max_row, column=1).font = FH
ws2.cell(row=ws2.max_row, column=5).font = FH
for col, w in zip("ABCDE", (12, 55, 26, 34, 15)):
    ws2.column_dimensions[col].width = w

ws3 = wb.create_sheet(f"A Pagar ({MES_PT[MES_SEG.month][:3]})")
cab(ws3, f"Contas a Pagar — {MES_PT[MES_SEG.month].capitalize()}/{MES_SEG.year}",
    ("Vencimento", "Descrição", "Conta", "Categoria", "Valor"))
for t in sorted(tx_apagar_seg, key=bdate):
    ws3.append((bdate(t), t.get("ds_transaction"), t.get("ds_account_main"), cat_raiz_nome(t), valor_efetivo(t) / 100))
ws3.append(("Total", "", "", "", tot_ap_seg / 100))
ws3.cell(row=ws3.max_row, column=1).font = FH
ws3.cell(row=ws3.max_row, column=5).font = FH
for col, w in zip("ABCDE", (12, 55, 26, 34, 15)):
    ws3.column_dimensions[col].width = w
wb.save(ARQ_XLSX)
print(f"OK: {ARQ_XLSX}")

if ENVIAR:
    KEY = open(os.path.join(_dir, "..", "relatorio_allt", ".resend_key")).read().strip()
    idem = "semanal-suporte-" + hashlib.md5((str(os.path.getmtime(ARQ_PDF)) + str(os.path.getmtime(ARQ_XLSX))).encode()).hexdigest()[:10]
    corpo = {
        "from": "Terceirizou <financeiro@terceirizou.com.br>",
        "to": ["henrique@terceirizou.com.br"],
        "bcc": ["financeiro@terceirizou.com.br"],
        "subject": f"Suporte Proteção Veicular — Relatório Semanal — {n_sexta}ª sexta de {MES_PT[MES_REF.month].capitalize()}/{MES_REF.year}",
        "html": (f"<p>Bom dia!</p><p>Segue o relatório semanal da <b>Suporte Proteção Veicular</b> — "
                 f"{n_sexta}ª sexta-feira de {MES_PT[MES_REF.month].capitalize()}/{MES_REF.year}.</p>"
                 f"<p>· <b>Consolidado pago</b> — 01/{MES_REF.month:02d} a {DIA_REF:02d}/{MES_REF.month:02d}: "
                 f"{brl(tot_e + tot_s)}<br>"
                 f"· <b>A pagar até {FIM_MES_REF.strftime('%d/%m')}</b>: {brl(tot_ap_mes)}<br>"
                 f"· <b>Contas a pagar de {MES_PT[MES_SEG.month]}</b>: {brl(tot_ap_seg)}</p>"
                 f"<p>Qualquer dúvida estamos à disposição.</p><p>Obrigado!</p>"),
        "attachments": [
            {"filename": os.path.basename(ARQ_PDF), "content": base64.b64encode(open(ARQ_PDF, "rb").read()).decode()},
            {"filename": os.path.basename(ARQ_XLSX), "content": base64.b64encode(open(ARQ_XLSX, "rb").read()).decode()},
        ],
    }
    UAH = {"Authorization": f"Bearer {KEY}", "User-Agent": "terceirizou-relatorios/1.0",
           "Content-Type": "application/json", "Idempotency-Key": idem}
    rq = urllib.request.Request("https://api.resend.com/emails", json.dumps(corpo).encode(), headers=UAH)
    with urllib.request.urlopen(rq, timeout=120) as resp:
        print("enviado:", json.loads(resp.read().decode()).get("id"))
