#!/usr/bin/env python3
# Relatórios DARKE ESTRATEGIA E NEGOCIOS LTDA — v1.0, 2026-10-08
#   Formato Campo Belo (logo capa + símbolo rodapé proporção real 372x553).
#   TUDO em REGIME DE CAIXA (situation in 1,2, mês por dt_billing) — pedido Vinícius 08/10.
#   Relatórios:
#     1. Comparativo dos últimos 13 meses por categoria (caixa)
#     2. Entradas e Saídas por Categoria do mês anterior (caixa)
#     3. Previsão de Entradas e Saídas por Categoria do mês corrente (caixa — pago + pendente)
#     4. Resultado dos últimos 12 meses por categoria (caixa)
#     5. Inadimplência até o último dia do mês anterior (receitas em aberto)
#     6. Previsão de Fluxo de Caixa 12m — lançamentos previstos por mês + saldo real;
#        a INADIMPLÊNCIA entra como previsão de receber 2 meses à frente (out → dez),
#        registrado na observação do relatório.
#   Filtro: exclusão SÓ categoria 99.01 (regra Vinícius 05/10 — descrição não filtra).
#   Cliente: consultoria estratégia e negócios (novo segmento). Contas: Sicoob Darke 213171,
#   Conta Carteira - Cobranças 228450, Conta Inicial 211966, Asaas 227034.
# Uso: python3 relatorio_darke.py [YYYY-MM-DD]  (default: hoje)
import json, os, sys, urllib.request, time
from collections import defaultdict
from datetime import date, timedelta

from reportlab.lib.pagesizes import A4
from reportlab.lib.units import cm
from reportlab.lib import colors
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.platypus import (SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle,
                                PageBreak, Image as RLImage, KeepTogether)

BASE = "https://api-v1.controlle.com"
_dir = os.path.dirname(os.path.abspath(__file__))
TOKEN = os.environ.get("CONTROLLE_TOKEN_DARKE") or open(os.path.join(_dir, ".controlle_token")).read().strip()
UA = {"Authorization": f"Bearer {TOKEN}", "User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) Chrome/126.0"}
LARANJA = colors.HexColor("#ff501c")
LARANJA_CLARO = colors.HexColor("#ffe3d6")
PRETO = colors.HexColor("#1a1a1a")
CINZA = colors.HexColor("#f5f5f5")
VERDE = colors.HexColor("#1a7f37")
VERMELHO = colors.HexColor("#c0392b")
LOGO = os.path.join(_dir, "logo-terceirizou.png")
SIMBOLO = os.path.join(_dir, "simbolo-terceirizou.png")

MES_PT = {1:"janeiro",2:"fevereiro",3:"março",4:"abril",5:"maio",6:"junho",7:"julho",8:"agosto",9:"setembro",10:"outubro",11:"novembro",12:"dezembro"}
MES_AB = {1:"jan",2:"fev",3:"mar",4:"abr",5:"mai",6:"jun",7:"jul",8:"ago",9:"set",10:"out",11:"nov",12:"dez"}

def req(url, tries=4):
    for i in range(tries):
        try:
            r = urllib.request.Request(url)
            for k, v in UA.items():
                r.add_header(k, v)
            with urllib.request.urlopen(r, timeout=120) as resp:
                return json.loads(resp.read().decode("utf-8", "replace"))
        except urllib.error.HTTPError as e:
            if e.code in (500, 502, 503) and i < tries - 1:
                time.sleep(3 * (i + 1))
                continue
            raise

def tx_list(start, end, **filtros):
    out, page = [], 1
    while True:
        url = (f"{BASE}/transaction/v1/transactions/list?start_date={start}&end_date={end}"
               f"&page={page}&orderBy=date&orderByCardinality=ASC")
        for k, v in filtros.items():
            url += f"&{k}={v}"
        tl = req(url).get("results", {}).get("transactionsList", [])
        out.extend(tl)
        if len(tl) < 100:
            return out
        page += 1

def brl(cents):
    v = cents / 100
    s = f"{abs(v):,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")
    return ("-" if v < 0 else "") + "R$ " + s

def brl_int(cents):
    v = int(round(cents / 100.0))
    s = f"{abs(v):,}".replace(",", ".")
    return ("-" if v < 0 else "") + s

def add_months(d, n):
    m = d.month - 1 + n
    y = d.year + m // 12
    m = m % 12 + 1
    return date(y, m, 1)

def bdate(t):
    return (t.get("dt_billing") or t.get("dt_due") or "")[:10]

def cat_nome(t):
    cats = t.get("apportionments_plan_account") or []
    return (cats[0].get("ds_category") or "?") if cats else "(sem categoria)"

def ok_ub(t):
    """Exclui SÓ a categoria 99.01 — descrição não filtra."""
    for c in (t.get("apportionments_plan_account") or []):
        if (c.get("ds_category") or "").startswith("99.01"):
            return False
    return True

# ===== datas =====
HOJE = date.fromisoformat(sys.argv[1]) if len(sys.argv) > 1 else date.today()
MES_ANT = add_months(HOJE.replace(day=1), -1)
FIM_MES_ANT = MES_ANT.replace(day=28) + timedelta(days=4)
FIM_MES_ANT = FIM_MES_ANT - timedelta(days=FIM_MES_ANT.day)
MES_COR = HOJE.replace(day=1)
FIM_MES_COR = MES_COR.replace(day=28) + timedelta(days=4)
FIM_MES_COR = FIM_MES_COR - timedelta(days=FIM_MES_COR.day)
FIM_PROJ = add_months(MES_COR, 11)
FIM_PROJ = FIM_PROJ.replace(day=28) + timedelta(days=4)
FIM_PROJ = FIM_PROJ - timedelta(days=FIM_PROJ.day)
MES_INAD = add_months(MES_COR, 2)  # inadimplência recebível 2 meses à frente

HOJE_LABEL = HOJE.strftime("%d/%m/%Y")
FIM_MES_ANT_LABEL = FIM_MES_ANT.strftime("%d/%m/%Y")

# ===== dados =====
# janela: 2 anos (13m do comparativo) + ano seguinte (projeção)
_tx = []
for y in sorted({add_months(MES_ANT, -12).year, MES_ANT.year}):
    _tx.extend(tx_list(f"{y}-01-01", f"{y}-12-31"))
_fut = tx_list(f"{MES_ANT.year + 1}-01-01", f"{MES_ANT.year + 1}-12-31")

def meses_janela(n, ini_base):
    out = []
    for i in range(n - 1, -1, -1):
        ini_m = add_months(ini_base, -i)
        fim_m = add_months(ini_m, 1) - timedelta(days=1)
        out.append((ini_m.isoformat(), fim_m.isoformat(), f"{MES_AB[ini_m.month]}/{str(ini_m.year)[2:]}"))
    return out
MESES13 = meses_janela(13, MES_ANT)

# 1. comparativo 13m (caixa)
matriz_13 = defaultdict(lambda: defaultdict(int))
for t in _tx:
    if not ok_ub(t) or t.get("situation") not in (1, 2): continue
    mes = bdate(t)[:7]
    if mes < MESES13[0][0][:7] or mes > MESES13[-1][1][:7]: continue
    aps = t.get("apportionments_plan_account") or []
    if len(aps) <= 1:
        matriz_13[cat_nome(t)][mes] += t["value_in_cent"]
    else:
        for c in aps:
            matriz_13[c.get("ds_category") or "?"][mes] += c.get("value") or 0
meses_com_mov = []
for _, fim_m_iso, lab in MESES13:
    tot = sum(v[fim_m_iso[:7]] for v in matriz_13.values())
    if tot != 0:
        meses_com_mov.append((fim_m_iso, lab))

# 2. mês anterior (caixa)
mes_ant_tx = [t for t in _tx if bdate(t)[:7] == f"{MES_ANT.year}-{MES_ANT.month:02d}"
              and t.get("situation") in (1, 2) and ok_ub(t)]
ant_rec = defaultdict(lambda: [0, 0])
ant_desp = defaultdict(lambda: [0, 0])
for t in mes_ant_tx:
    v = t["value_in_cent"]
    g = ant_rec if v > 0 else ant_desp
    g[cat_nome(t)][0] += v
    g[cat_nome(t)][1] += 1
ant_entradas = sum(v for v, _ in ant_rec.values())
ant_saidas = sum(v for v, _ in ant_desp.values())
ant_resultado = ant_entradas + ant_saidas

# 3. previsão do mês corrente (caixa — pago + pendente)
prev_cor = [t for t in _tx + _fut if MES_COR.isoformat() <= bdate(t) <= FIM_MES_COR.isoformat() and ok_ub(t)]
prev_rec = defaultdict(lambda: [0, 0])
prev_desp = defaultdict(lambda: [0, 0])
for t in prev_cor:
    v = t["value_in_cent"]
    g = prev_rec if v > 0 else prev_desp
    g[cat_nome(t)][0] += v
    g[cat_nome(t)][1] += 1
prev_entradas = sum(v for v, _ in prev_rec.values())
prev_saidas = sum(v for v, _ in prev_desp.values())

# 4. resultado 12m por categoria (caixa)
MESES12 = meses_janela(12, MES_ANT)
res12 = defaultdict(int)
for t in _tx:
    if not ok_ub(t) or t.get("situation") not in (1, 2): continue
    mes = bdate(t)[:7]
    if mes < MESES12[0][0][:7] or mes > MESES12[-1][1][:7]: continue
    aps = t.get("apportionments_plan_account") or []
    if len(aps) <= 1:
        res12[cat_nome(t)] += t["value_in_cent"]
    else:
        for c in aps:
            res12[c.get("ds_category") or "?"] += c.get("value") or 0

# 5. inadimplência até fim do mês anterior (receitas em aberto)
inad = [t for t in _tx if t["activity_type"] == 1 and t.get("situation") == 0
        and bdate(t) <= FIM_MES_ANT.isoformat() and ok_ub(t)]
g_inad = defaultdict(lambda: [0, 0])
for t in inad:
    g_inad[cat_nome(t)][0] += t["value_in_cent"]
    g_inad[cat_nome(t)][1] += 1
total_inad = sum(v for v, _ in g_inad.values())

# 6. projeção 12m (lançamentos previstos por mês + saldo real + inadimplência 2m à frente)
all_fut = _tx + _fut
proj = []
for i in range(12):
    ini_m = add_months(MES_COR, i)
    fim_m = add_months(ini_m, 1) - timedelta(days=1)
    lab = f"{MES_AB[ini_m.month]}/{str(ini_m.year)[2:]}"
    txs_m = [t for t in all_fut if ini_m.isoformat() <= bdate(t) <= fim_m.isoformat() and ok_ub(t)]
    ent = sum(t["value_in_cent"] for t in txs_m if t["value_in_cent"] > 0)
    sai = sum(t["value_in_cent"] for t in txs_m if t["value_in_cent"] < 0)
    inad_mes = total_inad if ini_m == MES_INAD else 0
    proj.append((lab, ent, sai, inad_mes))

# saldos das contas no último dia do mês anterior
contas_api = [c for c in req(f"{BASE}/account/v1/accounts").get("results", []) if c.get("status") == 1]
saldos_conta = []
for c in contas_api:
    try:
        b = req(f"{BASE}/transaction/v1/transactions/balances?start_date=2017-01-01&end_date={FIM_MES_ANT.isoformat()}&id_account_main={c['id']}")["results"]
        saldos_conta.append((c["ds_account"], b["balanceDone"]))
    except Exception:
        pass
saldos_conta = [(n, v) for n, v in saldos_conta if v != 0]
total_saldos = sum(v for _, v in saldos_conta)
saldo_ac = total_saldos
proj_ac = []
for lab, ent, sai, inad_mes in proj:
    saldo_ac += ent + sai + inad_mes
    proj_ac.append((lab, ent, sai, inad_mes, saldo_ac))

print("dados prontos")

# ===== PDF =====
styles = getSampleStyleSheet()
h1 = ParagraphStyle("h1", parent=styles["Heading1"], fontSize=16, textColor=PRETO, spaceAfter=4)
h2 = ParagraphStyle("h2", parent=styles["Heading2"], fontSize=12, textColor=LARANJA, spaceBefore=14, spaceAfter=3)
h3 = ParagraphStyle("h3", parent=styles["Heading3"], fontSize=10, textColor=PRETO, spaceBefore=10, spaceAfter=3)
sub = ParagraphStyle("sub", parent=styles["Normal"], fontName="Helvetica-Oblique", fontSize=8.5, textColor=colors.HexColor("#777777"), spaceAfter=10)
cell = ParagraphStyle("cell", parent=styles["Normal"], fontSize=8)
cellb = ParagraphStyle("cellb", parent=styles["Normal"], fontSize=8, fontName="Helvetica-Bold")
cellr = ParagraphStyle("cellr", parent=cell, alignment=2)
cellrb = ParagraphStyle("cellrb", parent=cellb, alignment=2)
cellc = ParagraphStyle("cellc", parent=cell, alignment=1)
P = Paragraph

def P_val(cents, style):
    s = ParagraphStyle("v", parent=style, textColor=(VERDE if cents > 0 else VERMELHO if cents < 0 else PRETO))
    return P(brl(cents), s)

def P_val_int(cents, style):
    s = ParagraphStyle("v", parent=style, textColor=(VERDE if cents > 0 else VERMELHO if cents < 0 else PRETO))
    return P(brl_int(cents), s)

def tabela(rows, widths, fs=7.5):
    t = Table(rows, colWidths=widths, repeatRows=1)
    style = [("FONTNAME", (0,0), (-1,-1), "Helvetica"), ("FONTSIZE", (0,0), (-1,-1), fs),
             ("TOPPADDING", (0,0), (-1,-1), 3), ("BOTTOMPADDING", (0,0), (-1,-1), 3),
             ("VALIGN", (0,0), (-1,-1), "MIDDLE"),
             ("BACKGROUND", (0,0), (-1,0), LARANJA), ("TEXTCOLOR", (0,0), (-1,0), colors.white),
             ("FONTNAME", (0,0), (-1,0), "Helvetica-Bold"),
             ("LINEBELOW", (0,-1), (-1,-1), 0.7, LARANJA)]
    for i in range(1, len(rows)):
        if i % 2 == 0:
            style.append(("BACKGROUND", (0,i), (-1,i), CINZA))
    t.setStyle(TableStyle(style))
    return t

def tabela_grupo(grupos, titulo, total_label):
    rows = [[P("<b>Categoria</b>", cell), P("<b>Lançamentos</b>", cellc), P("<b>Valor</b>", cellr)]]
    for nome, (v, n) in sorted(grupos.items(), key=lambda x: x[0]):
        rows.append([P(nome, cell), P(str(n), cellc), P_val(v, cellr)])
    rows.append([P(f"<b>{total_label}</b>", cellrb), P(f"<b>{sum(n for _, n in grupos.values())}</b>", cellc), P_val(sum(v for v, _ in grupos.values()), cellrb)])
    t = tabela(rows, [11*cm, 2.5*cm, 3*cm])
    t.setStyle(TableStyle([("BACKGROUND", (0,len(rows)-1), (-1,len(rows)-1), LARANJA_CLARO)]))
    return [P(f"<b>{titulo}</b>", h2), t]

def rodape_simbolo(canvas, doc):
    if os.path.exists(SIMBOLO):
        alt = 1.2 * cm
        canvas.drawImage(SIMBOLO, 18.4*cm, 1.1*cm, width=alt*372/553, height=alt, mask="auto")

ARQ_PDF = f"artifacts/{HOJE.strftime('%y%m%d')}_Relatorio_Darke.pdf"
doc = SimpleDocTemplate(ARQ_PDF, pagesize=A4, leftMargin=1.5*cm, rightMargin=1.5*cm,
                        topMargin=1.3*cm, bottomMargin=1.6*cm,
                        title="Relatório Gerencial — Darke Estratégia e Negócios")
E = []
h1c = ParagraphStyle("h1c", parent=h1, fontSize=13, alignment=1, spaceAfter=1)
h2c = ParagraphStyle("h2c", parent=h2, fontSize=11, alignment=1, spaceBefore=2, spaceAfter=2)
subc = ParagraphStyle("subc", parent=sub, fontSize=8, alignment=1, spaceAfter=0)
if os.path.exists(LOGO):
    img = RLImage(LOGO, width=5.5*cm, height=5.5*cm*561/1600)
    img.hAlign = "CENTER"
    E.append(img)
E.append(Spacer(1, 8))
E.append(P("<b>DARKE ESTRATÉGIA E NEGÓCIOS</b>", h1c))
E.append(P("Relatório Gerencial Mensal", h2c))
E.append(P(f"Gerado em {HOJE_LABEL} · Fonte: Controlle · Ref.: {MES_PT[MES_ANT.month].capitalize()} de {MES_ANT.year} — regime de caixa", subc))
E.append(Spacer(1, 6))

# 1. comparativo 13m (receitas primeiro, despesas depois)
E.append(P(f"Comparativo dos Últimos 13 Meses por Categoria ({MESES13[0][2]} a {MESES13[-1][2]}) — regime de caixa", h2))
E.append(P("Meses sem movimentação foram excluídos do comparativo.", sub))
cat_names = sorted({c for c in matriz_13}, key=lambda c: sum(matriz_13[c].values()), reverse=True)
n_col = len(meses_com_mov)
w_cat = max(4.5, 16.5 - 1.3 * n_col - 2.8)
cm_rows = [[P("<b>Categoria</b>", cell)] + [P(f"<b>{lab}</b>", cellr) for _, lab in meses_com_mov] + [P("<b>Média</b>", cellr), P("<b>Total</b>", cellr)]]
for cat in cat_names:
    row = [P(cat, cell)]
    vals = [matriz_13[cat].get(fim_m_iso[:7], 0) for fim_m_iso, _ in meses_com_mov]
    for v in vals:
        row.append(P_val_int(v, cellr) if v else P("—", cellr))
    row.append(P_val_int(round(sum(vals) / len(vals)), cellrb))
    row.append(P_val_int(sum(vals), cellrb))
    cm_rows.append(row)
res_row = [P("<b>Resultado do mês</b>", cellb)]
for fim_m_iso, _ in meses_com_mov:
    tot_mes = sum(v[fim_m_iso[:7]] for v in matriz_13.values())
    res_row.append(P_val_int(tot_mes, cellrb))
_res_mensais = [sum(v[fim_m_iso[:7]] for v in matriz_13.values()) for fim_m_iso, _ in meses_com_mov]
res_row.append(P_val_int(round(sum(_res_mensais) / len(_res_mensais)), cellrb))
res_row.append(P_val_int(sum(_res_mensais), cellrb))
cm_rows.append(res_row)
if n_col <= 7:
    E.append(tabela(cm_rows, [w_cat*cm] + [1.3*cm]*n_col + [1.4*cm, 1.4*cm], fs=6.5))
else:
    E.append(tabela(cm_rows, [4.5*cm] + [1.05*cm]*n_col + [1.2*cm, 1.2*cm], fs=5.8))

# 2. mês anterior
E.append(PageBreak())
E += tabela_grupo(ant_rec, f"Entradas — {MES_PT[MES_ANT.month].capitalize()} de {MES_ANT.year} (regime de caixa)", "Total de Entradas")
E += tabela_grupo(ant_desp, f"Saídas — {MES_PT[MES_ANT.month].capitalize()} de {MES_ANT.year} (regime de caixa)", "Total de Saídas")
res_rows = [[P("<b>Resultado do mês</b>", cellrb), P(f"<b>{len(mes_ant_tx)}</b>", cellc), P_val(ant_resultado, cellrb)]]
t = tabela(res_rows, [11*cm, 2.5*cm, 3*cm])
t.setStyle(TableStyle([("BACKGROUND", (0,0), (-1,0), LARANJA_CLARO)]))
E.append(t)

# 3. previsão do mês corrente
E.append(PageBreak())
E += tabela_grupo(prev_rec, f"Previsão de Entradas — {MES_PT[MES_COR.month].capitalize()} de {MES_COR.year} (pago + pendente)", "Total de Entradas Previstas")
E += tabela_grupo(prev_desp, f"Previsão de Saídas — {MES_PT[MES_COR.month].capitalize()} de {MES_COR.year} (pago + pendente)", "Total de Saídas Previstas")
res_rows = [[P("<b>Resultado previsto do mês</b>", cellrb), P(f"<b>{len(prev_cor)}</b>", cellc), P_val(prev_entradas + prev_saidas, cellrb)]]
t = tabela(res_rows, [11*cm, 2.5*cm, 3*cm])
t.setStyle(TableStyle([("BACKGROUND", (0,0), (-1,0), LARANJA_CLARO)]))
E.append(t)

# 4. resultado 12m por categoria
E.append(PageBreak())
r12_rows = [[P("<b>Categoria</b>", cell), P("<b>Resultado 12 meses</b>", cellr)]]
for cat in sorted(res12, key=lambda c: res12[c], reverse=True):
    r12_rows.append([P(cat, cell), P_val(res12[cat], cellr)])
r12_rows.append([P("<b>Resultado total 12 meses</b>", cellrb), P_val(sum(res12.values()), cellrb)])
t = tabela(r12_rows, [11*cm, 5*cm])
t.setStyle(TableStyle([("BACKGROUND", (0,len(r12_rows)-1), (-1,len(r12_rows)-1), LARANJA_CLARO)]))
E.append(KeepTogether([P(f"Resultado dos Últimos 12 Meses por Categoria ({MESES12[0][2]} a {MESES12[-1][2]}) — regime de caixa", h2), t]))

# 5. inadimplência
E.append(Spacer(1, 14))
inad_rows = [[P("<b>Categoria</b>", cell), P("<b>Receitas em aberto</b>", cellr)]]
for cat, (v, n) in sorted(g_inad.items(), key=lambda x: x[0]):
    inad_rows.append([P(f"{cat} ({n} lanç.)", cell), P_val(v, cellr)])
inad_rows.append([P("<b>Total em aberto</b>", cellrb), P_val(total_inad, cellrb)])
t = tabela(inad_rows, [11*cm, 5*cm])
t.setStyle(TableStyle([("BACKGROUND", (0,len(inad_rows)-1), (-1,len(inad_rows)-1), LARANJA_CLARO)]))
E.append(KeepTogether([P(f"Inadimplência até {FIM_MES_ANT_LABEL}", h2), t]))
if inad:
    rows = [[P("<b>Vencimento</b>", cell), P("<b>Descrição</b>", cell), P("<b>Conta</b>", cell), P("<b>Valor</b>", cellr)]]
    for t in sorted(inad, key=bdate):
        d = bdate(t)
        rows.append([P(f"{d[8:10]}/{d[5:7]}/{d[:4]}", cell), P((t.get("ds_transaction") or "")[:70], cell),
                     P(t.get("ds_account_main") or "", cell), P_val(t["value_in_cent"], cellr)])
    tt = tabela(rows, [2.2*cm, 8.8*cm, 3.5*cm, 3*cm])
    tt.setStyle(TableStyle([("BACKGROUND", (0,len(rows)-1), (-1,len(rows)-1), LARANJA_CLARO)]))
    E.append(KeepTogether([P("<b>Demonstrativo</b>", h3), tt]))

# 6. projeção 12m
E.append(PageBreak())
rows6 = [[P("<b>Mês</b>", cell), P("<b>Entradas previstas</b>", cellr), P("<b>Saídas previstas</b>", cellr),
          P("<b>Inadimplência a receber</b>", cellr), P("<b>Saldo previsto</b>", cellr)]]
for lab, ent, sai, inad_mes, saldo in proj_ac:
    rows6.append([P(lab, cell), P_val(ent, cellr), P_val(sai, cellr),
                  (P_val(inad_mes, cellr) if inad_mes else P("—", cellr)), P_val(saldo, cellrb)])
t = tabela(rows6, [2.6*cm, 3.6*cm, 3.6*cm, 3.6*cm, 3.6*cm], fs=7.5)
E.append(P(f"Previsão de Fluxo de Caixa — Próximos 12 Meses ({MES_PT[MES_COR.month].capitalize()} de {MES_COR.year} a {MES_PT[FIM_PROJ.month].capitalize()} de {FIM_PROJ.year})", h2))
E.append(t)
E.append(P(f"Observação: a inadimplência de {brl(total_inad)} (receitas em aberto até {FIM_MES_ANT_LABEL}) "
           f"entra como previsão de receber em {MES_PT[MES_INAD.month].capitalize()} de {MES_INAD.year} "
           f"(dois meses à frente do relatório). Saldo de partida = saldo real em {FIM_MES_ANT_LABEL}; "
           "movimentos = lançamentos previstos do Controlle.", sub))

E.append(Spacer(1, 10))
E.append(P(f"Saldo nas contas em {FIM_MES_ANT_LABEL}", h2))
sc_rows = [[P("<b>Conta</b>", cell), P(f"<b>Saldo em {FIM_MES_ANT_LABEL}</b>", cellr)]]
for nome, v in saldos_conta:
    sc_rows.append([P(nome, cell), P_val(v, cellr)])
sc_rows.append([P("<b>Total</b>", cellrb), P_val(total_saldos, cellrb)])
t = tabela(sc_rows, [11*cm, 5*cm])
t.setStyle(TableStyle([("BACKGROUND", (0,len(sc_rows)-1), (-1,len(sc_rows)-1), LARANJA_CLARO)]))
E.append(t)

E.append(Spacer(1, 10))
E.append(P("Gerado automaticamente pela Terceirizou · dados do Controlle", sub))
doc.build(E, onFirstPage=rodape_simbolo, onLaterPages=rodape_simbolo)
print(f"OK: {ARQ_PDF}")

# ===== Excel =====
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

ARQ_XLSX = f"artifacts/{HOJE.strftime('%y%m%d')}_Relatorio_Darke.xlsx"
wb = Workbook()
wb.remove(wb.active)
FILL_H = PatternFill("solid", fgColor="FF501C")
FILL_T = PatternFill("solid", fgColor="FFE3D6")
FH = Font(bold=True, color="FFFFFF")
FB = Font(bold=True)
VERDE_XL = "1A7F37"
VERMELHO_XL = "C0392B"
TOT_LABELS = ("Total", "Resultado")

def aba(nome, linhas, larguras, titulo=None):
    ws = wb.create_sheet(nome[:31])
    if titulo:
        ws.append((titulo,))
        ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=max(1, len(larguras)))
        c = ws.cell(row=1, column=1)
        c.font = Font(bold=True, size=12, color="FF501C")
        c.alignment = Alignment(horizontal="center", vertical="center")
        ws.row_dimensions[1].height = 22
        desloc = 1
    else:
        desloc = 0
    for r in linhas:
        ws.append(list(r))
    hdr_row = desloc + 1
    headers = [str(ws.cell(row=hdr_row, column=j).value or "") for j in range(1, ws.max_column + 1)]
    for c in ws[hdr_row]:
        c.fill = FILL_H
        c.font = FH
        c.alignment = Alignment(horizontal="center", vertical="center")
    for row in ws.iter_rows(min_row=hdr_row + 1):
        label = str(row[0].value or "")
        is_total = label.startswith(TOT_LABELS)
        if is_total:
            for c in row:
                c.fill = FILL_T
                c.alignment = Alignment(horizontal="center", vertical="center")
            row[0].font = FB
        for j, c in enumerate(row[1:], 2):
            if isinstance(c.value, (int, float)):
                hdr = headers[j-1] if j-1 < len(headers) else ""
                if "Lançamentos" in hdr:
                    c.number_format = "0"
                    c.font = Font(bold=is_total)
                else:
                    c.number_format = '"R$" #,##0.00'
                    cor = VERDE_XL if c.value > 0 else (VERMELHO_XL if c.value < 0 else "1A1A1A")
                    c.font = Font(bold=is_total, color=cor)
    for j, w in enumerate(larguras, 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = ws.cell(row=hdr_row + 1, column=1)

r_ = lambda v: v / 100

# 1. comparativo 13m
aba("Comparativo 13m",
    [("Categoria",) + tuple(lab for _, lab in meses_com_mov) + ("Média", "Total")] +
    [(cat,) + tuple((r_(matriz_13[cat].get(fim_m_iso[:7], 0)) if matriz_13[cat].get(fim_m_iso[:7], 0) else None) for fim_m_iso, _ in meses_com_mov)
     + (r_(round(sum(matriz_13[cat].values()) / len(meses_com_mov))), r_(sum(matriz_13[cat].values())),)
     for cat in cat_names] +
    [("Resultado do mês",) + tuple(int(round(m / 100)) if m else None for m in _res_mensais)
     + (int(round(sum(_res_mensais) / len(_res_mensais) / 100)), int(round(sum(_res_mensais) / 100)))],
    [40] + [13]*n_col + [13, 15], titulo=f"Comparativo dos Últimos 13 Meses por Categoria (regime de caixa)")

# 2. mês anterior
aba("Ent. Saídas " + MES_AB[MES_ANT.month],
    [("Categoria", "Lançamentos", "Valor")] +
    [(n, n2, r_(v)) for n, (v, n2) in sorted(ant_rec.items())] +
    [("Total de Entradas", sum(n for _, n in ant_rec.values()), r_(ant_entradas))] +
    [(n, n2, r_(v)) for n, (v, n2) in sorted(ant_desp.items())] +
    [("Total de Saídas", sum(n for _, n in ant_desp.values()), r_(ant_saidas)),
     ("Resultado do mês", len(mes_ant_tx), r_(ant_resultado))],
    [50, 14, 16], titulo=f"Entradas e Saídas — {MES_PT[MES_ANT.month].capitalize()} de {MES_ANT.year} (regime de caixa)")

# 3. previsão do mês corrente
aba("Previsão " + MES_AB[MES_COR.month],
    [("Categoria", "Lançamentos", "Valor")] +
    [(n, n2, r_(v)) for n, (v, n2) in sorted(prev_rec.items())] +
    [("Total de Entradas Previstas", sum(n for _, n in prev_rec.values()), r_(prev_entradas))] +
    [(n, n2, r_(v)) for n, (v, n2) in sorted(prev_desp.items())] +
    [("Total de Saídas Previstas", sum(n for _, n in prev_desp.values()), r_(prev_saidas)),
     ("Resultado previsto do mês", len(prev_cor), r_(prev_entradas + prev_saidas))],
    [50, 14, 16], titulo=f"Previsão de Entradas e Saídas — {MES_PT[MES_COR.month].capitalize()} de {MES_COR.year} (pago + pendente)")

# 4. resultado 12m
aba("Resultado 12m",
    [("Categoria", "Resultado 12 meses")] +
    [(cat, r_(res12[cat])) for cat in sorted(res12, key=lambda c: res12[c], reverse=True)] +
    [("Resultado total 12 meses", r_(sum(res12.values())))],
    [50, 20], titulo=f"Resultado dos Últimos 12 Meses por Categoria ({MESES12[0][2]} a {MESES12[-1][2]}) — regime de caixa")

# 5. inadimplência
aba("Inadimplência",
    [("Categoria", "Receitas em aberto")] +
    [(f"{cat} ({n} lanç.)", r_(v)) for cat, (v, n) in sorted(g_inad.items(), key=lambda x: x[0])] +
    [("Total em aberto", r_(total_inad))] +
    [("", "")] +
    [("Vencimento", "Descrição")] +
    [(f"{bdate(t)[8:10]}/{bdate(t)[5:7]}/{bdate(t)[:4]}", (t.get("ds_transaction") or "")[:70]) for t in sorted(inad, key=bdate)],
    [50, 60], titulo=f"Inadimplência até {FIM_MES_ANT_LABEL}")

# 6. projeção 12m
aba("Proj fluxo 12m",
    [("Mês", "Entradas previstas", "Saídas previstas", "Inadimplência a receber", "Saldo previsto")] +
    [(lab, r_(ent), r_(sai), (r_(inad_mes) if inad_mes else None), r_(saldo)) for lab, ent, sai, inad_mes, saldo in proj_ac],
    [10, 18, 18, 22, 18], titulo=f"Previsão de Fluxo de Caixa — Próximos 12 Meses (inadimplência: receber em {MES_AB[MES_INAD.month]}/{str(MES_INAD.year)[2:]})")

# saldos
aba("Saldo contas",
    [("Conta", f"Saldo em {FIM_MES_ANT_LABEL}")] +
    [(n, r_(v)) for n, v in saldos_conta] + [("Total", r_(total_saldos))],
    [30, 18], titulo=f"Saldo nas contas em {FIM_MES_ANT_LABEL}")

wb.save(ARQ_XLSX)
print(f"OK: {ARQ_XLSX}")
