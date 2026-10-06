#!/usr/bin/env python3
# Relatórios CARNIEL EMPREENDIMENTO LTDA — v1.0, 2026-10-06
#   Formato Campo Belo (logo capa compacta + símbolo no rodapé com proporção real 372x553).
#   Relatórios (pedido Vinícius 06/10, regime de CAIXA — situation in 1,2, mês por dt_billing):
#     1. Entradas e Saídas por categoria SINTÉTICO (mês anterior) — 1 linha por categoria, sem lançamentos
#     2. Entradas e Saídas por categoria ANALÍTICO — categorias + TODOS os lançamentos de cada uma
#     3. Saldo nas contas no último dia do mês anterior
#   Filtro: exclusão SÓ categoria 99.01 (regra Vinícius 05/10 — descrição não filtra).
#   Cliente: incorporadora/construção civil (novo segmento) — 8 contas, categorias próprias de obra.
# Uso: python3 relatorio_carniel.py [YYYY-MM-DD]  (default: hoje)
import json, os, sys, urllib.request, time
from collections import defaultdict
from datetime import date, timedelta

from reportlab.lib.pagesizes import A4
from reportlab.lib.units import cm
from reportlab.lib import colors
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.platypus import (SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle,
                                PageBreak, Image as RLImage, KeepTogether)

BASE = "https://api-v1.controlle.com"
_dir = os.path.dirname(os.path.abspath(__file__))
TOKEN = os.environ.get("CONTROLLE_TOKEN_CARNIEL") or open(os.path.join(_dir, ".controlle_token")).read().strip()
UA = {"Authorization": f"Bearer {TOKEN}", "User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) Chrome/126.0"}
LARANJA = colors.HexColor("#ff501c")
LARANJA_CLARO = colors.HexColor("#ffe3d6")
PRETO = colors.HexColor("#1a1a1a")
CINZA = colors.HexColor("#f5f5f5")
VERDE = colors.HexColor("#1a7f37")
VERMELHO = colors.HexColor("#c0392b")
LOGO = os.path.join(_dir, "logo-terceirizou.png")
SIMBOLO = os.path.join(_dir, "simbolo-terceirizou.png")

MES_PT = {1:"janeiro",2:"fevereiro",3:"março",4:"abril",5:"maio",6:"junho",7:"julho",8:"agosto",9:"setembro",10:"outubro",11:"novembro",12:"dezembro"}
MES_AB = {1:"jan",2:"fev",3:"mar",4:"abr",5:"mai",6:"jun",7:"jul",8:"ago",9:"set",10:"out",11:"nov",12:"dez"}

def req(url, tries=4):
    for i in range(tries):
        try:
            r = urllib.request.Request(url)
            for k, v in UA.items():
                r.add_header(k, v)
            with urllib.request.urlopen(r, timeout=120) as resp:
                return json.loads(resp.read().decode("utf-8", "replace"))
        except urllib.error.HTTPError as e:
            if e.code in (500, 502, 503) and i < tries - 1:
                time.sleep(3 * (i + 1))
                continue
            raise

def tx_list(start, end, **filtros):
    out, page = [], 1
    while True:
        url = (f"{BASE}/transaction/v1/transactions/list?start_date={start}&end_date={end}"
               f"&page={page}&orderBy=date&orderByCardinality=ASC")
        for k, v in filtros.items():
            url += f"&{k}={v}"
        tl = req(url).get("results", {}).get("transactionsList", [])
        out.extend(tl)
        if len(tl) < 100:
            return out
        page += 1

def brl(cents):
    v = cents / 100
    s = f"{abs(v):,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")
    return ("-" if v < 0 else "") + "R$ " + s

def add_months(d, n):
    m = d.month - 1 + n
    y = d.year + m // 12
    m = m % 12 + 1
    return date(y, m, 1)

def bdate(t):
    return (t.get("dt_billing") or t.get("dt_due") or "")[:10]

def cat_nome(t):
    cats = t.get("apportionments_plan_account") or []
    return (cats[0].get("ds_category") or "?") if cats else "(sem categoria)"

def ok_ub(t):
    """Exclui SÓ a categoria 99.01 (regra Vinícius 05/10) — descrição não filtra."""
    for c in (t.get("apportionments_plan_account") or []):
        if (c.get("ds_category") or "").startswith("99.01"):
            return False
    return True

# ===== datas =====
HOJE = date.fromisoformat(sys.argv[1]) if len(sys.argv) > 1 else date.today()
MES_ANT = add_months(HOJE.replace(day=1), -1)
FIM_MES_ANT = MES_ANT.replace(day=28) + timedelta(days=4)
FIM_MES_ANT = FIM_MES_ANT - timedelta(days=FIM_MES_ANT.day)

HOJE_LABEL = HOJE.strftime("%d/%m/%Y")
FIM_MES_ANT_LABEL = FIM_MES_ANT.strftime("%d/%m/%Y")

# ===== dados =====
_tx = tx_list(f"{MES_ANT.year}-01-01", f"{MES_ANT.year}-12-31")
mes_ant_tx = [t for t in _tx if bdate(t)[:7] == f"{MES_ANT.year}-{MES_ANT.month:02d}"
              and t.get("situation") in (1, 2) and ok_ub(t)]

# agrupamento por categoria (receitas separadas das despesas)
rec = defaultdict(lambda: [0, 0])
desp = defaultdict(lambda: [0, 0])
for t in mes_ant_tx:
    v = t["value_in_cent"]
    g = rec if v > 0 else desp
    g[cat_nome(t)][0] += v
    g[cat_nome(t)][1] += 1
tot_rec = sum(v for v, _ in rec.values())
tot_desp = sum(v for v, _ in desp.values())
resultado = tot_rec + tot_desp

# saldos das contas no último dia do mês anterior
contas_api = [c for c in req(f"{BASE}/account/v1/accounts").get("results", []) if c.get("status") == 1]
saldos_conta = []
for c in contas_api:
    try:
        b = req(f"{BASE}/transaction/v1/transactions/balances?start_date=2017-01-01&end_date={FIM_MES_ANT.isoformat()}&id_account_main={c['id']}")["results"]
        saldos_conta.append((c["ds_account"], b["balanceDone"]))
    except Exception:
        pass
saldos_conta = [(n, v) for n, v in saldos_conta if v != 0]
total_saldos = sum(v for _, v in saldos_conta)

print("dados prontos")

# ===== PDF =====
styles = getSampleStyleSheet()
h1 = ParagraphStyle("h1", parent=styles["Heading1"], fontSize=16, textColor=PRETO, spaceAfter=4)
h2 = ParagraphStyle("h2", parent=styles["Heading2"], fontSize=12, textColor=LARANJA, spaceBefore=14, spaceAfter=3)
h3 = ParagraphStyle("h3", parent=styles["Heading3"], fontSize=10, textColor=PRETO, spaceBefore=10, spaceAfter=3)
sub = ParagraphStyle("sub", parent=styles["Normal"], fontName="Helvetica-Oblique", fontSize=8.5, textColor=colors.HexColor("#777777"), spaceAfter=10)
cell = ParagraphStyle("cell", parent=styles["Normal"], fontSize=8)
cellb = ParagraphStyle("cellb", parent=styles["Normal"], fontSize=8, fontName="Helvetica-Bold")
cellr = ParagraphStyle("cellr", parent=cell, alignment=2)
cellrb = ParagraphStyle("cellrb", parent=cellb, alignment=2)
cellc = ParagraphStyle("cellc", parent=cell, alignment=1)
P = Paragraph

def P_val(cents, style):
    s = ParagraphStyle("v", parent=style, textColor=(VERDE if cents > 0 else VERMELHO if cents < 0 else PRETO))
    return P(brl(cents), s)

def tabela(rows, widths, fs=7.5):
    t = Table(rows, colWidths=widths, repeatRows=1)
    style = [("FONTNAME", (0,0), (-1,-1), "Helvetica"), ("FONTSIZE", (0,0), (-1,-1), fs),
             ("TOPPADDING", (0,0), (-1,-1), 3), ("BOTTOMPADDING", (0,0), (-1,-1), 3),
             ("VALIGN", (0,0), (-1,-1), "MIDDLE"),
             ("BACKGROUND", (0,0), (-1,0), LARANJA), ("TEXTCOLOR", (0,0), (-1,0), colors.white),
             ("FONTNAME", (0,0), (-1,0), "Helvetica-Bold"),
             ("LINEBELOW", (0,-1), (-1,-1), 0.7, LARANJA)]
    for i in range(1, len(rows)):
        if i % 2 == 0:
            style.append(("BACKGROUND", (0,i), (-1,i), CINZA))
    t.setStyle(TableStyle(style))
    return t

def rodape_simbolo(canvas, doc):
    if os.path.exists(SIMBOLO):
        alt = 1.2 * cm
        canvas.drawImage(SIMBOLO, 18.4*cm, 1.1*cm, width=alt*372/553, height=alt, mask="auto")

ARQ_PDF = f"artifacts/{HOJE.strftime('%y%m%d')}_Relatorio_Carniel.pdf"
doc = SimpleDocTemplate(ARQ_PDF, pagesize=A4, leftMargin=1.5*cm, rightMargin=1.5*cm,
                        topMargin=1.3*cm, bottomMargin=1.6*cm,
                        title="Relatório Gerencial — Carniel Empreendimento")
E = []
h1c = ParagraphStyle("h1c", parent=h1, fontSize=13, alignment=1, spaceAfter=1)
h2c = ParagraphStyle("h2c", parent=h2, fontSize=11, alignment=1, spaceBefore=2, spaceAfter=2)
subc = ParagraphStyle("subc", parent=sub, fontSize=8, alignment=1, spaceAfter=0)
if os.path.exists(LOGO):
    img = RLImage(LOGO, width=5.5*cm, height=5.5*cm*561/1600)
    img.hAlign = "CENTER"
    E.append(img)
E.append(Spacer(1, 8))
E.append(P("<b>CARNIEL EMPREENDIMENTO</b>", h1c))
E.append(P("Relatório Gerencial Mensal", h2c))
E.append(P(f"Gerado em {HOJE_LABEL} · Fonte: Controlle · Ref.: {MES_PT[MES_ANT.month].capitalize()} de {MES_ANT.year} — regime de caixa", subc))
E.append(Spacer(1, 6))

# 1. SINTÉTICO — 1 linha por categoria (receitas primeiro, despesas depois)
def tabela_grupo(grupos, titulo, total_label):
    rows = [[P("<b>Categoria</b>", cell), P("<b>Lançamentos</b>", cellc), P("<b>Valor</b>", cellr)]]
    for nome, (v, n) in sorted(grupos.items(), key=lambda x: x[0]):
        rows.append([P(nome, cell), P(str(n), cellc), P_val(v, cellr)])
    rows.append([P(f"<b>{total_label}</b>", cellrb), P(f"<b>{sum(n for _, n in grupos.values())}</b>", cellc), P_val(sum(v for v, _ in grupos.values()), cellrb)])
    t = tabela(rows, [11*cm, 2.5*cm, 3*cm])
    t.setStyle(TableStyle([("BACKGROUND", (0,len(rows)-1), (-1,len(rows)-1), LARANJA_CLARO)]))
    return [P(f"<b>{titulo}</b>", h2), t]

bloco1 = [P(f"Entradas e Saídas por Categoria — Sintético — {MES_PT[MES_ANT.month].capitalize()} de {MES_ANT.year} (regime de caixa)", h2)]
bloco1 += tabela_grupo(rec, "Entradas", "Total de Entradas")
bloco1 += tabela_grupo(desp, "Saídas", "Total de Saídas")
res_rows = [[P("<b>Resultado do mês</b>", cellrb), P(f"<b>{len(mes_ant_tx)}</b>", cellc), P_val(resultado, cellrb)]]
t = tabela(res_rows, [11*cm, 2.5*cm, 3*cm])
t.setStyle(TableStyle([("BACKGROUND", (0,0), (-1,0), LARANJA_CLARO)]))
bloco1.append(t)
E.append(KeepTogether(bloco1))

# 2. ANALÍTICO — categorias + todos os lançamentos
E.append(PageBreak())
E.append(P(f"Entradas e Saídas por Categoria — Analítico — {MES_PT[MES_ANT.month].capitalize()} de {MES_ANT.year} (regime de caixa)", h2))
E += tabela_grupo(rec, "Entradas", "Total de Entradas")
E += tabela_grupo(desp, "Saídas", "Total de Saídas")
res_rows = [[P("<b>Resultado do mês</b>", cellrb), P(f"<b>{len(mes_ant_tx)}</b>", cellc), P_val(resultado, cellrb)]]
t = tabela(res_rows, [11*cm, 2.5*cm, 3*cm])
t.setStyle(TableStyle([("BACKGROUND", (0,0), (-1,0), LARANJA_CLARO)]))
E.append(t)
E.append(Spacer(1, 14))
E.append(P("<b>Lançamentos por categoria</b>", h2))
por_cat = defaultdict(list)
for t in mes_ant_tx:
    por_cat[cat_nome(t)].append(t)
for cat in sorted(por_cat):
    txs_cat = sorted(por_cat[cat], key=bdate)
    rows = [[P("<b>Data</b>", cell), P("<b>Descrição</b>", cell), P("<b>Conta</b>", cell), P("<b>Valor</b>", cellr)]]
    for t in txs_cat:
        d = bdate(t)
        rows.append([P(f"{d[8:10]}/{d[5:7]}/{d[:4]}", cell), P((t.get("ds_transaction") or "")[:70], cell),
                     P(t.get("ds_account_main") or "", cell), P_val(t["value_in_cent"], cellr)])
    rows.append([P("<b>Total</b>", cellrb), P(f"<b>{cat}</b>", cellrb), P(f"<b>{len(txs_cat)}</b>", cellc), P_val(sum(t["value_in_cent"] for t in txs_cat), cellrb)])
    tt = tabela(rows, [2.2*cm, 8.3*cm, 3.5*cm, 3*cm])
    tt.setStyle(TableStyle([("BACKGROUND", (0,len(rows)-1), (-1,len(rows)-1), LARANJA_CLARO)]))
    E.append(KeepTogether([P(f"<b>{cat}</b>", h3), tt]))

# 3. saldo nas contas
E.append(PageBreak())
sc_rows = [[P("<b>Conta</b>", cell), P(f"<b>Saldo em {FIM_MES_ANT_LABEL}</b>", cellr)]]
for nome, v in saldos_conta:
    sc_rows.append([P(nome, cell), P_val(v, cellr)])
sc_rows.append([P("<b>Total</b>", cellrb), P_val(total_saldos, cellrb)])
t = tabela(sc_rows, [11*cm, 5*cm])
t.setStyle(TableStyle([("BACKGROUND", (0,len(sc_rows)-1), (-1,len(sc_rows)-1), LARANJA_CLARO)]))
E.append(KeepTogether([P(f"Saldo nas contas em {FIM_MES_ANT_LABEL}", h2), t]))

E.append(Spacer(1, 10))
E.append(P("Gerado automaticamente pela Terceirizou · dados do Controlle", sub))
doc.build(E, onFirstPage=rodape_simbolo, onLaterPages=rodape_simbolo)
print(f"OK: {ARQ_PDF}")

# ===== Excel =====
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

ARQ_XLSX = f"artifacts/{HOJE.strftime('%y%m%d')}_Relatorio_Carniel.xlsx"
wb = Workbook()
wb.remove(wb.active)
FILL_H = PatternFill("solid", fgColor="FF501C")
FILL_T = PatternFill("solid", fgColor="FFE3D6")
FH = Font(bold=True, color="FFFFFF")
FB = Font(bold=True)
VERDE_XL = "1A7F37"
VERMELHO_XL = "C0392B"
TOT_LABELS = ("Total", "Resultado")

def aba(nome, linhas, larguras, titulo=None):
    ws = wb.create_sheet(nome[:31])
    if titulo:
        ws.append((titulo,))
        ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=max(1, len(larguras)))
        c = ws.cell(row=1, column=1)
        c.font = Font(bold=True, size=12, color="FF501C")
        c.alignment = Alignment(horizontal="center", vertical="center")
        ws.row_dimensions[1].height = 22
        desloc = 1
    else:
        desloc = 0
    for r in linhas:
        ws.append(list(r))
    hdr_row = desloc + 1
    headers = [str(ws.cell(row=hdr_row, column=j).value or "") for j in range(1, ws.max_column + 1)]
    for c in ws[hdr_row]:
        c.fill = FILL_H
        c.font = FH
        c.alignment = Alignment(horizontal="center", vertical="center")
    for row in ws.iter_rows(min_row=hdr_row + 1):
        label = str(row[0].value or "")
        is_total = label.startswith(TOT_LABELS)
        if is_total:
            for c in row:
                c.fill = FILL_T
                c.alignment = Alignment(horizontal="center", vertical="center")
            row[0].font = FB
        for j, c in enumerate(row[1:], 2):
            if isinstance(c.value, (int, float)):
                hdr = headers[j-1] if j-1 < len(headers) else ""
                if "Lançamentos" in hdr:
                    c.number_format = "0"
                    c.font = Font(bold=is_total)
                else:
                    c.number_format = '"R$" #,##0.00'
                    cor = VERDE_XL if c.value > 0 else (VERMELHO_XL if c.value < 0 else "1A1A1A")
                    c.font = Font(bold=is_total, color=cor)
    for j, w in enumerate(larguras, 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = ws.cell(row=hdr_row + 1, column=1)

r_ = lambda v: v / 100
fmt_d = lambda d: f"{d[8:10]}/{d[5:7]}/{d[:4]}"

def linhas_grupo(grupos, total_label):
    out = [("Categoria", "Lançamentos", "Valor")]
    for nome, (v, n) in sorted(grupos.items(), key=lambda x: x[0]):
        out.append((nome, n, r_(v)))
    out.append((total_label, sum(n for _, n in grupos.values()), r_(sum(v for v, _ in grupos.values()))))
    return out

# 1. sintético
aba("Sintético",
    linhas_grupo(rec, "Total de Entradas")[0:1] + linhas_grupo(rec, "Total de Entradas")[1:] +
    linhas_grupo(desp, "Total de Saídas")[1:] +
    [("Resultado do mês", len(mes_ant_tx), r_(resultado))],
    [50, 14, 16], titulo=f"Entradas e Saídas por Categoria — Sintético — {MES_PT[MES_ANT.month].capitalize()} de {MES_ANT.year} (regime de caixa)")

# 2. analítico (categorias + lançamentos)
por_cat = defaultdict(list)
for t in mes_ant_tx:
    por_cat[cat_nome(t)].append(t)
an_rows = [("Categoria", "Data", "Descrição", "Conta", "Valor")]
for cat in sorted(por_cat):
    for t in sorted(por_cat[cat], key=bdate):
        an_rows.append((cat, fmt_d(bdate(t)), (t.get("ds_transaction") or "")[:80],
                        t.get("ds_account_main") or "", r_(t["value_in_cent"])))
    an_rows.append((f"Total {cat}", "", "", "", r_(sum(t["value_in_cent"] for t in por_cat[cat]))))
an_rows.append(("Resultado do mês", f"{len(mes_ant_tx)} lançamentos", "", "", r_(resultado)))
aba("Analítico", an_rows, [40, 12, 58, 24, 16],
    titulo=f"Entradas e Saídas por Categoria — Analítico — {MES_PT[MES_ANT.month].capitalize()} de {MES_ANT.year} (regime de caixa)")

# 3. saldos
aba("Saldo contas",
    [("Conta", f"Saldo em {FIM_MES_ANT_LABEL}")] +
    [(n, r_(v)) for n, v in saldos_conta] + [("Total", r_(total_saldos))],
    [30, 18], titulo=f"Saldo nas contas em {FIM_MES_ANT_LABEL}")

wb.save(ARQ_XLSX)
print(f"OK: {ARQ_XLSX}")
