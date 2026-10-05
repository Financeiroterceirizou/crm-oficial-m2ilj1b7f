#!/usr/bin/env python3
# Relatório Mensal — GIANE ZANELATO SAUDE CAPILAR LTDA — v1.2, 2026-10-05
#   v1.2 (feedback Vinícius): filtro de exclusão = SÓ a categoria 99.01 Transferência entre Contas.
#     A descrição NÃO filtra — nesta licença os Pix de clientes chegam como "Transferência <nome>"
#     com categoria de receita (1.01/1.02/1.03) e os pagamentos por Pix como "Transferência <fornecedor>"
#     com categoria de despesa (6.01/1.01) — o filtro antigo por descrição engolia R$ 14.302,70 de
#     receitas e R$ 9.533,95 de despesas reais de setembro.
#   v1.1 (feedback Vinícius): comparativo 13m e Receitas/Despesas do mês anterior em REGIME DE CAIXA
#     (situation in 1,2 = pago + agendado, mês por dt_billing); Previsão de Receitas e Despesas
#     substituída por Previsão de DESPESAS do mês corrente (caixa); título da Projeção de Fluxo
#     de Caixa 6m movido para ACIMA da tabela.
# Uso: python3 relatorio_giane_zanelato.py [YYYY-MM-DD]  (default: hoje)
# Gera UM PDF + UM Excel. Relatórios:
#   1. Comparativo dos últimos 13 meses por categoria (regime de caixa) — EXCLUI colunas de meses
#      sem movimentação (pedido: "quando não tiver movimentação excluir a coluna do mês")
#   2. Receitas e Despesas do mês anterior (regime de caixa)
#   3. Despesas em aberto até o último dia do mês anterior
#   4. Previsão de Despesas do mês corrente (pago + pendente)
#   5. Previsão de Fluxo de Caixa para os próximos 06 meses (lançamentos previstos + saldo real)
# Fonte: API Controlle v1. Conta dedicada, sem centro de custo.
# Contas: Sicoob 226898, Conta Inicial 226623 (0), Nubank Giane 226900, Nubank Nicolas 226901, CAIXINHA 227915.
# Categorias PRÓPRIAS (1.01 Receitas com Tratamento, 6.01 Despesas Pessoais dos Sócios etc.) —
#   não usa o plano de contas das vistorias. Movimentação começa em jun/jul/26.
# Envio: vinicius@terceirizou.com.br (padrão; destino do cliente a definir pelo Vinícius).
import json, os, sys, urllib.request, time
from collections import defaultdict
from datetime import date, timedelta

from reportlab.lib.pagesizes import A4
from reportlab.lib.units import cm
from reportlab.lib import colors
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.platypus import (SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle,
                                PageBreak, Image as RLImage)

BASE = "https://api-v1.controlle.com"
_dir = os.path.dirname(os.path.abspath(__file__))
TOKEN = os.environ.get("CONTROLLE_TOKEN_GIANE") or open(os.path.join(_dir, ".controlle_token")).read().strip()
UA = {"Authorization": f"Bearer {TOKEN}", "User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) Chrome/126.0"}
LARANJA = colors.HexColor("#ff501c")
LARANJA_CLARO = colors.HexColor("#ffe3d6")
PRETO = colors.HexColor("#1a1a1a")
CINZA = colors.HexColor("#f5f5f5")
VERDE = colors.HexColor("#1a7f37")
VERMELHO = colors.HexColor("#c0392b")
LOGO = os.path.join(_dir, "logo-terceirizou.png")

MES_PT = {1:"janeiro",2:"fevereiro",3:"março",4:"abril",5:"maio",6:"junho",7:"julho",8:"agosto",9:"setembro",10:"outubro",11:"novembro",12:"dezembro"}
MES_AB = {1:"jan",2:"fev",3:"mar",4:"abr",5:"mai",6:"jun",7:"jul",8:"ago",9:"set",10:"out",11:"nov",12:"dez"}

def req(url, tries=4):
    for i in range(tries):
        try:
            r = urllib.request.Request(url)
            for k, v in UA.items():
                r.add_header(k, v)
            with urllib.request.urlopen(r, timeout=120) as resp:
                return json.loads(resp.read().decode("utf-8", "replace"))
        except urllib.error.HTTPError as e:
            if e.code in (500, 502, 503) and i < tries - 1:
                time.sleep(3 * (i + 1))
                continue
            raise

def tx_list(start, end, **filtros):
    out, page = [], 1
    while True:
        url = (f"{BASE}/transaction/v1/transactions/list?start_date={start}&end_date={end}"
               f"&page={page}&orderBy=date&orderByCardinality=ASC")
        for k, v in filtros.items():
            url += f"&{k}={v}"
        tl = req(url).get("results", {}).get("transactionsList", [])
        out.extend(tl)
        if len(tl) < 100:
            return out
        page += 1

def brl(cents):
    v = cents / 100
    s = f"{abs(v):,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")
    return ("-" if v < 0 else "") + "R$ " + s

def brl_int(cents):
    v = int(round(cents / 100.0))
    s = f"{abs(v):,}".replace(",", ".")
    return ("-" if v < 0 else "") + s

def add_months(d, n):
    m = d.month - 1 + n
    y = d.year + m // 12
    m = m % 12 + 1
    return date(y, m, 1)

def bdate(t):
    return (t.get("dt_billing") or t.get("dt_due") or "")[:10]

def cat_nome(t):
    cats = t.get("apportionments_plan_account") or []
    return (cats[0].get("ds_category") or "?") if cats else "(sem categoria)"

def ok_ub(t):
    """Exclui SÓ a categoria 99.01 Transferência entre Contas (regra Vinícius 05/10).
    A descrição NÃO filtra — "Transferência <nome>" com categoria de receita é Pix de cliente."""
    for c in (t.get("apportionments_plan_account") or []):
        if (c.get("ds_category") or "").startswith("99.01"):
            return False
    return True

def agrupa_por_categoria(txs, so_negativas=False, so_positivas=False):
    g = defaultdict(lambda: [0, 0])
    for t in txs:
        v = t["value_in_cent"]
        if so_negativas and v >= 0: continue
        if so_positivas and v <= 0: continue
        g[cat_nome(t)][0] += v
        g[cat_nome(t)][1] += 1
    return g

# ===== datas =====
HOJE = date.fromisoformat(sys.argv[1]) if len(sys.argv) > 1 else date.today()
MES_ANT = add_months(HOJE.replace(day=1), -1)
FIM_MES_ANT = MES_ANT.replace(day=28) + timedelta(days=4)
FIM_MES_ANT = FIM_MES_ANT - timedelta(days=FIM_MES_ANT.day)
MES_COR = HOJE.replace(day=1)
FIM_MES_COR = MES_COR.replace(day=28) + timedelta(days=4)
FIM_MES_COR = FIM_MES_COR - timedelta(days=FIM_MES_COR.day)
INI_13 = add_months(MES_ANT, -12)   # 13 meses

HOJE_LABEL = HOJE.strftime("%d/%m/%Y")
FIM_MES_ANT_LABEL = FIM_MES_ANT.strftime("%d/%m/%Y")

# ===== dados =====
# comparativo + mês anterior: 2 anos (INI_13.year + MES_ANT.year); projeção: ano seguinte (_fut)
_tx = []
for y in (INI_13.year, MES_ANT.year):
    _tx.extend(tx_list(f"{y}-01-01", f"{y}-12-31"))
# dedup por id (defensivo)
_vistos = set()
_tx_dedup = []
for t in _tx:
    tid = t.get("id_transactions")
    if tid and tid in _vistos:
        continue
    if tid:
        _vistos.add(tid)
    _tx_dedup.append(t)
_tx = _tx_dedup

def meses_13():
    out = []
    for i in range(0, 13):
        ini_m = add_months(INI_13, i)
        fim_m = add_months(ini_m, 1) - timedelta(days=1)
        out.append((ini_m.isoformat(), fim_m.isoformat(), f"{MES_AB[ini_m.month]}/{str(ini_m.year)[2:]}"))
    return out
MESES13 = meses_13()

# 1. Comparativo 13 meses — REGIME DE CAIXA (situation in 1,2 = pago + agendado, mês por dt_billing)
matriz_13 = defaultdict(lambda: defaultdict(int))
for t in _tx:
    if not ok_ub(t):
        continue
    if t.get("situation") not in (1, 2):
        continue
    mes = bdate(t)[:7]
    if mes < MESES13[0][0][:7] or mes > MESES13[-1][1][:7]:
        continue
    aps = t.get("apportionments_plan_account") or []
    if len(aps) <= 1:
        matriz_13[cat_nome(t)][mes] += t["value_in_cent"]
    else:
        for c in aps:
            matriz_13[c.get("ds_category") or "?"][mes] += c.get("value") or 0
# meses COM movimentação (colunas do comparativo — pedido: excluir meses sem movimentação)
meses_com_mov = []
for _, fim_m_iso, lab in MESES13:
    tot = sum(v[fim_m_iso[:7]] for v in matriz_13.values())
    if tot != 0:
        meses_com_mov.append((fim_m_iso, lab))

# 2. Receitas e Despesas do mês anterior — REGIME DE CAIXA (situation in 1,2 + dt_billing)
mes_ant_tx = [t for t in _tx if bdate(t)[:7] == f"{MES_ANT.year}-{MES_ANT.month:02d}"
              and t.get("situation") in (1, 2) and ok_ub(t)]
ant_rec = agrupa_por_categoria(mes_ant_tx, so_positivas=True)
ant_desp = agrupa_por_categoria(mes_ant_tx, so_negativas=True)
ant_entradas = sum(v for v, _ in ant_rec.values())
ant_saidas = sum(v for v, _ in ant_desp.values())
ant_resultado = ant_entradas + ant_saidas

# 3. Despesas em aberto até o último dia do mês anterior
desp_aberto = [t for t in _tx if t["activity_type"] == 0 and t.get("situation") == 0
               and bdate(t) <= FIM_MES_ANT.isoformat() and ok_ub(t)]
g_desp_aberto = agrupa_por_categoria(desp_aberto)
total_desp_aberto = sum(v for v, _ in g_desp_aberto.values())

# 4. Previsão de Despesas do mês corrente (pago + pendente)
prev_cor = [t for t in _tx if MES_COR.isoformat() <= bdate(t) <= FIM_MES_COR.isoformat() and ok_ub(t)]
prev_desp = agrupa_por_categoria(prev_cor, so_negativas=True)
prev_desp_total = sum(v for v, _ in prev_desp.values())

# 5. Previsão de Fluxo de Caixa — próximos 6 meses (lançamentos previstos por mês + saldo real)
_fut = []
try:
    _fut = tx_list(f"{MES_ANT.year + 1}-01-01", f"{MES_ANT.year + 1}-12-31")
except Exception:
    _fut = []
proj = []
for i in range(1, 7):
    ini_m = add_months(MES_COR, i - 1)
    fim_m = add_months(ini_m, 1) - timedelta(days=1)
    lab = f"{MES_AB[ini_m.month]}/{str(ini_m.year)[2:]}"
    txs_m = [t for t in (_tx + _fut) if ini_m.isoformat() <= bdate(t) <= fim_m.isoformat() and ok_ub(t)]
    ent = sum(t["value_in_cent"] for t in txs_m if t["activity_type"] == 1)
    sai = sum(t["value_in_cent"] for t in txs_m if t["activity_type"] == 0)
    proj.append((lab, ent, sai))

# saldos das contas no último dia do mês anterior
contas = [c for c in req(f"{BASE}/account/v1/accounts").get("results", []) if c.get("status") == 1]
saldos_conta = []
for c in contas:
    try:
        b = req(f"{BASE}/transaction/v1/transactions/balances?start_date=2017-01-01&end_date={FIM_MES_ANT.isoformat()}&id_account_main={c['id']}")["results"]
        saldos_conta.append((c["ds_account"], b["balanceDone"]))
    except Exception:
        pass
saldos_conta = [(n, v) for n, v in saldos_conta if v != 0]
total_saldos = sum(v for _, v in saldos_conta)
# projeção com saldo acumulado (partida = saldo real 30/09)
saldo_ac = total_saldos
proj_ac = []
for lab, ent, sai in proj:
    saldo_ac += ent + sai
    proj_ac.append((lab, ent, sai, saldo_ac))

print("dados prontos")

# ===== PDF =====
styles = getSampleStyleSheet()
h1 = ParagraphStyle("h1", parent=styles["Heading1"], fontSize=16, textColor=PRETO, spaceAfter=4)
h2 = ParagraphStyle("h2", parent=styles["Heading2"], fontSize=12, textColor=LARANJA, spaceBefore=14, spaceAfter=3)
h3 = ParagraphStyle("h3", parent=styles["Heading3"], fontSize=10, textColor=PRETO, spaceBefore=10, spaceAfter=3)
sub = ParagraphStyle("sub", parent=styles["Normal"], fontName="Helvetica-Oblique", fontSize=8.5, textColor=colors.HexColor("#777777"), spaceAfter=10)
cell = ParagraphStyle("cell", parent=styles["Normal"], fontSize=8)
cellb = ParagraphStyle("cellb", parent=styles["Normal"], fontSize=8, fontName="Helvetica-Bold")
cellr = ParagraphStyle("cellr", parent=cell, alignment=2)
cellrb = ParagraphStyle("cellrb", parent=cellb, alignment=2)
cellc = ParagraphStyle("cellc", parent=cell, alignment=1)

def P_val(cents, style):
    s = ParagraphStyle("v", parent=style, textColor=(VERDE if cents > 0 else VERMELHO if cents < 0 else PRETO))
    return P(brl(cents), s)

def P_val_int(cents, style):
    s = ParagraphStyle("v", parent=style, textColor=(VERDE if cents > 0 else VERMELHO if cents < 0 else PRETO))
    return P(brl_int(cents), s)

def tabela(rows, widths, fs=7.5):
    t = Table(rows, colWidths=widths, repeatRows=1)
    style = [("FONTNAME", (0,0), (-1,-1), "Helvetica"), ("FONTSIZE", (0,0), (-1,-1), fs),
             ("TOPPADDING", (0,0), (-1,-1), 3), ("BOTTOMPADDING", (0,0), (-1,-1), 3),
             ("VALIGN", (0,0), (-1,-1), "MIDDLE"),
             ("BACKGROUND", (0,0), (-1,0), LARANJA), ("TEXTCOLOR", (0,0), (-1,0), colors.white),
             ("FONTNAME", (0,0), (-1,0), "Helvetica-Bold"),
             ("LINEBELOW", (0,-1), (-1,-1), 0.7, LARANJA)]
    for i in range(1, len(rows)):
        if i % 2 == 0:
            style.append(("BACKGROUND", (0,i), (-1,i), CINZA))
    t.setStyle(TableStyle(style))
    return t

def tabela_cat(titulo, grupos, total_label="Total"):
    rows = [[P("<b>Categoria</b>", cell), P("<b>Lançamentos</b>", cellc), P("<b>Valor</b>", cellr)]]
    for nome, (v, n) in sorted(grupos.items(), key=lambda x: x[0]):
        rows.append([P(nome, cell), P(str(n), cellc), P_val(v, cellr)])
    tot = sum(v for v, _ in grupos.values())
    rows.append([P(f"<b>{total_label}</b>", cellrb), P(f"<b>{sum(n for _, n in grupos.values())}</b>", cellc), P_val(tot, cellrb)])
    t = tabela(rows, [11*cm, 2.5*cm, 3*cm])
    t.setStyle(TableStyle([("BACKGROUND", (0,len(rows)-1), (-1,len(rows)-1), LARANJA_CLARO)]))
    return [P(f"<b>{titulo}</b>", h2), t]

P = Paragraph
ARQ_PDF = f"artifacts/{HOJE.strftime('%y%m%d')}_Relatorio_Giane_Zanelato.pdf"
doc = SimpleDocTemplate(ARQ_PDF, pagesize=A4, leftMargin=1.5*cm, rightMargin=1.5*cm, topMargin=1.3*cm, bottomMargin=1.3*cm,
                        title="Relatório Gerencial — Giane Zanelato Saúde Capilar")
E = []

# capa compacta + 1º relatório na MESMA página (padrão dos relatórios)
h1c = ParagraphStyle("h1c", parent=h1, fontSize=13, alignment=1, spaceAfter=1)
h2c = ParagraphStyle("h2c", parent=h2, fontSize=11, alignment=1, spaceBefore=2, spaceAfter=2)
subc = ParagraphStyle("subc", parent=sub, fontSize=8, alignment=1, spaceAfter=0)
if os.path.exists(LOGO):
    img = RLImage(LOGO, width=5.5*cm, height=5.5*cm*561/1600)
    img.hAlign = "CENTER"
    E.append(img)
E.append(Spacer(1, 10))
E.append(P("<b>GIANE ZANELATO SAÚDE CAPILAR</b>", h1c))
E.append(P("Relatório Gerencial Mensal", h2c))
E.append(P(f"Gerado em {HOJE_LABEL} · Fonte: Controlle · Ref.: {MES_PT[MES_ANT.month].capitalize()} de {MES_ANT.year}", subc))
E.append(Spacer(1, 8))

# 1. Comparativo 13 meses (só meses COM movimentação)
E.append(P(f"Comparativo dos Últimos 13 Meses por Categoria ({MES_AB[INI_13.month]}/{str(INI_13.year)[2:]} a {MES_AB[MES_ANT.month]}/{str(MES_ANT.year)[2:]}) — regime de caixa", h2))
E.append(P("Meses sem movimentação foram excluídos do comparativo.", sub))
cat_names = sorted({c for c in matriz_13})
n_col = len(meses_com_mov)
w_cat = max(4.5, 16.5 - 1.3 * n_col - 2.8)
cm_rows = [[P("<b>Categoria</b>", cell)] + [P(f"<b>{lab}</b>", cellr) for _, lab in meses_com_mov] + [P("<b>Média</b>", cellr), P("<b>Total</b>", cellr)]]
for cat in cat_names:
    row = [P(cat, cell)]
    vals = [matriz_13[cat].get(fim_m_iso[:7], 0) for fim_m_iso, _ in meses_com_mov]
    for v in vals:
        row.append(P_val_int(v, cellr) if v else P("—", cellr))
    row.append(P_val_int(round(sum(vals) / len(vals)), cellrb))
    row.append(P_val_int(sum(vals), cellrb))
    cm_rows.append(row)
# linha de resultado do mês
res_row = [P("<b>Resultado do mês</b>", cellb)]
for fim_m_iso, _ in meses_com_mov:
    tot_mes = sum(v[fim_m_iso[:7]] for v in matriz_13.values())
    res_row.append(P_val_int(tot_mes, cellrb))
_res_mensais = [sum(v[fim_m_iso[:7]] for v in matriz_13.values()) for fim_m_iso, _ in meses_com_mov]
res_row.append(P_val_int(round(sum(_res_mensais) / len(_res_mensais)), cellrb))
res_row.append(P_val_int(sum(_res_mensais), cellrb))
cm_rows.append(res_row)
if n_col <= 7:
    E.append(tabela(cm_rows, [w_cat*cm] + [1.3*cm]*n_col + [1.4*cm, 1.4*cm], fs=6.5))
else:
    E.append(tabela(cm_rows, [4.5*cm] + [1.05*cm]*n_col + [1.2*cm, 1.2*cm], fs=5.8))

# 2. Receitas e Despesas do mês anterior
E.append(PageBreak())
E += tabela_cat(f"Receitas e Despesas — {MES_PT[MES_ANT.month].capitalize()} de {MES_ANT.year} (regime de caixa)", ant_rec, total_label="Total de Receitas")
rd_rows = [[P("<b>Categoria</b>", cell), P("<b>Lançamentos</b>", cellc), P("<b>Valor</b>", cellr)]]
for nome, (v, n) in sorted(ant_desp.items()):
    rd_rows.append([P(nome, cell), P(str(n), cellc), P_val(v, cellr)])
rd_rows.append([P("<b>Total de Despesas</b>", cellrb), P(f"<b>{sum(n for _, n in ant_desp.values())}</b>", cellc), P_val(ant_saidas, cellrb)])
rd_rows.append([P("<b>Resultado do mês</b>", cellrb), P(f"<b>{len(mes_ant_tx)}</b>", cellc), P_val(ant_resultado, cellrb)])
t = tabela(rd_rows, [11*cm, 2.5*cm, 3*cm])
t.setStyle(TableStyle([("BACKGROUND", (0,len(rd_rows)-2), (-1,len(rd_rows)-1), LARANJA_CLARO)]))
E.append(t)

# 3. Despesas em aberto até o último dia do mês anterior
E.append(Spacer(1, 14))
E.append(P(f"Despesas em aberto até {FIM_MES_ANT_LABEL}", h2))
E += tabela_cat("Resumo por categoria", g_desp_aberto, total_label="Total em aberto")
if desp_aberto:
    por_cat = defaultdict(list)
    for t in desp_aberto:
        por_cat[cat_nome(t)].append(t)
    E.append(P("<b>Demonstrativo dos lançamentos</b>", h2))
    for cat in sorted(por_cat):
        txs_cat = sorted(por_cat[cat], key=bdate)
        rows = [[P("<b>Vencimento</b>", cell), P("<b>Descrição</b>", cell), P("<b>Conta</b>", cell), P("<b>Situação</b>", cellc), P("<b>Valor</b>", cellr)]]
        for t in txs_cat:
            d = bdate(t)
            rows.append([P(f"{d[8:10]}/{d[5:7]}/{d[:4]}", cell), P((t.get("ds_transaction") or "")[:70], cell),
                         P(t.get("ds_account_main") or "", cell), P("Aberto", cellc), P_val(t["value_in_cent"], cellr)])
        rows.append([P("<b>Total</b>", cellrb), P(f"<b>{cat}</b>", cellrb), P("", cell), P(f"<b>{len(txs_cat)}</b>", cellc), P_val(sum(t["value_in_cent"] for t in txs_cat), cellrb)])
        tt = tabela(rows, [2.2*cm, 7.3*cm, 3.2*cm, 1.8*cm, 3*cm])
        tt.setStyle(TableStyle([("BACKGROUND", (0,len(rows)-1), (-1,len(rows)-1), LARANJA_CLARO)]))
        E.append(P(f"<b>{cat}</b>", h3))
        E.append(tt)

# 4. Previsão de Despesas do mês corrente
E.append(PageBreak())
rows3 = [[P("<b>Categoria</b>", cell), P("<b>Lançamentos</b>", cellc), P("<b>Valor</b>", cellr)]]
for nm, (v, n) in sorted(prev_desp.items()):
    rows3.append([P(nm, cell), P(str(n), cellc), P_val(v, cellr)])
rows3.append([P("<b>Total de Despesas Previstas</b>", cellrb), P(f"<b>{sum(n for _, n in prev_desp.values())}</b>", cellc), P_val(prev_desp_total, cellrb)])
t = tabela(rows3, [11*cm, 2.5*cm, 3*cm])
t.setStyle(TableStyle([("BACKGROUND", (0,len(rows3)-1), (-1,len(rows3)-1), LARANJA_CLARO)]))
E.append(P(f"Previsão de Despesas — {MES_PT[MES_COR.month].capitalize()} de {MES_COR.year} (pago + pendente)", h2))
E.append(t)

# 5. Previsão de Fluxo de Caixa — próximos 6 meses (título ACIMA da tabela)
E.append(Spacer(1, 14))
E.append(P(f"Previsão de Fluxo de Caixa — Próximos 6 Meses ({MES_PT[MES_COR.month].capitalize()} de {MES_COR.year} a {MES_PT[add_months(MES_COR,5).month].capitalize()} de {add_months(MES_COR,5).year})", h2))
rows6 = [[P("<b>Mês</b>", cell), P("<b>Entradas previstas</b>", cellr), P("<b>Saídas previstas</b>", cellr), P("<b>Saldo previsto</b>", cellr)]]
for lab, ent, sai, saldo in proj_ac:
    rows6.append([P(lab, cell), P_val(ent, cellr), P_val(sai, cellr), P_val(saldo, cellrb)])
t = tabela(rows6, [3.2*cm, 4.4*cm, 4.4*cm, 4.6*cm], fs=7.5)
E.append(t)
E.append(P("Nota: saldo de partida = saldo real em " + FIM_MES_ANT_LABEL + "; movimentos = lançamentos previstos do Controlle. "
           "As receitas da clínica não são recorrentes no sistema — o saldo projetado considera apenas as despesas previstas; "
           "cada mês de vendas real reduz a queda.", sub))

E.append(Spacer(1, 10))
E.append(P(f"Saldo nas contas em {FIM_MES_ANT_LABEL}", h2))
sc_rows = [[P("<b>Conta</b>", cell), P(f"<b>Saldo em {FIM_MES_ANT_LABEL}</b>", cellr)]]
for nome, v in saldos_conta:
    sc_rows.append([P(nome, cell), P_val(v, cellr)])
sc_rows.append([P("<b>Total</b>", cellrb), P_val(total_saldos, cellrb)])
t = tabela(sc_rows, [11*cm, 5*cm])
t.setStyle(TableStyle([("BACKGROUND", (0,len(sc_rows)-1), (-1,len(sc_rows)-1), LARANJA_CLARO)]))
E.append(t)

E.append(Spacer(1, 10))
E.append(P("Gerado automaticamente pela Terceirizou · dados do Controlle", sub))

doc.build(E)
print(f"OK: {ARQ_PDF}")

# ===== Excel =====
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

ARQ_XLSX = f"artifacts/{HOJE.strftime('%y%m%d')}_Relatorio_Giane_Zanelato.xlsx"
wb = Workbook()
wb.remove(wb.active)
FILL_H = PatternFill("solid", fgColor="FF501C")
FILL_T = PatternFill("solid", fgColor="FFE3D6")
FH = Font(bold=True, color="FFFFFF")
FB = Font(bold=True)
VERDE_XL = "1A7F37"
VERMELHO_XL = "C0392B"
TOT_LABELS = ("Total", "Totais", "Resultado", "Previsão", "Total de Receitas", "Total de Despesas",
              "Total de Despesas Previstas", "Total previsto", "Total em aberto")

def aba(nome, linhas, larguras, titulo=None):
    ws = wb.create_sheet(nome[:31])
    if titulo:
        ws.append((titulo,))
        ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=max(1, len(larguras)))
        c = ws.cell(row=1, column=1)
        c.font = Font(bold=True, size=12, color="FF501C")
        c.alignment = Alignment(horizontal="center", vertical="center")
        ws.row_dimensions[1].height = 22
        desloc = 1
    else:
        desloc = 0
    for r in linhas:
        ws.append(list(r))
    hdr_row = desloc + 1
    headers = [str(ws.cell(row=hdr_row, column=j).value or "") for j in range(1, ws.max_column + 1)]
    for c in ws[hdr_row]:
        c.fill = FILL_H
        c.font = FH
        c.alignment = Alignment(horizontal="center", vertical="center")
    for row in ws.iter_rows(min_row=hdr_row + 1):
        label = str(row[0].value or "")
        is_total = label.startswith(TOT_LABELS)
        if is_total:
            for c in row:
                c.fill = FILL_T
                c.alignment = Alignment(horizontal="center", vertical="center")
            row[0].font = FB
        for j, c in enumerate(row[1:], 2):
            if isinstance(c.value, (int, float)):
                hdr = headers[j-1] if j-1 < len(headers) else ""
                if "Lançamentos" in hdr:
                    c.number_format = "0"
                    c.font = Font(bold=is_total)
                else:
                    c.number_format = '"R$" #,##0.00'
                    cor = VERDE_XL if c.value > 0 else (VERMELHO_XL if c.value < 0 else "1A1A1A")
                    c.font = Font(bold=is_total, color=cor)
    for j, w in enumerate(larguras, 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = ws.cell(row=hdr_row + 1, column=1)

r_ = lambda v: v / 100

# 1. comparativo 13m (só meses com movimentação)
aba("Comparativo 13m",
    [("Categoria",) + tuple(lab for _, lab in meses_com_mov) + ("Média", "Total")] +
    [(cat,) + tuple((r_(matriz_13[cat].get(fim_m_iso[:7], 0)) if matriz_13[cat].get(fim_m_iso[:7], 0) else None) for fim_m_iso, _ in meses_com_mov)
     + (r_(round(sum(matriz_13[cat].values()) / len(meses_com_mov))), r_(sum(matriz_13[cat].values())),)
     for cat in cat_names] +
    [("Resultado do mês",) + tuple(int(round(m / 100)) if m else None for m in _res_mensais)
     + (int(round(sum(_res_mensais) / len(_res_mensais) / 100)), int(round(sum(_res_mensais) / 100)))],
    [40] + [13]*n_col + [13, 15], titulo=f"Comparativo dos Últimos 13 Meses por Categoria (regime de caixa) — meses sem movimentação excluídos")

# 2. mês anterior
aba("Receitas e Despesas " + MES_AB[MES_ANT.month],
    [("Categoria", "Lançamentos", "Valor")] +
    [(n, n2, r_(v)) for n, (v, n2) in sorted(ant_rec.items())] +
    [("Total de Receitas", sum(n for _, n in ant_rec.values()), r_(ant_entradas))] +
    [(n, n2, r_(v)) for n, (v, n2) in sorted(ant_desp.items())] +
    [("Total de Despesas", sum(n for _, n in ant_desp.values()), r_(ant_saidas)),
     ("Resultado do mês", len(mes_ant_tx), r_(ant_resultado))],
    [45, 14, 16], titulo=f"Receitas e Despesas — {MES_PT[MES_ANT.month].capitalize()} de {MES_ANT.year} (regime de caixa)")

# 3. despesas em aberto
aba("Despesas em aberto",
    [("Categoria", "Lançamentos", "Valor")] +
    [(n, n2, r_(v)) for n, (v, n2) in sorted(g_desp_aberto.items())] +
    [("Total em aberto", len(desp_aberto), r_(total_desp_aberto))],
    [45, 14, 16], titulo=f"Despesas em aberto até {FIM_MES_ANT_LABEL}")

# 4. previsão de despesas do mês corrente
aba("Previsão Despesas " + MES_AB[MES_COR.month],
    [("Categoria", "Lançamentos", "Valor")] +
    [(n, n2, r_(v)) for n, (v, n2) in sorted(prev_desp.items())] +
    [("Total de Despesas Previstas", sum(n for _, n in prev_desp.values()), r_(prev_desp_total))],
    [45, 14, 16], titulo=f"Previsão de Despesas — {MES_PT[MES_COR.month].capitalize()} de {MES_COR.year} (pago + pendente)")

# 5. projeção 6m
aba("Proj fluxo 6m",
    [("Mês", "Entradas previstas", "Saídas previstas", "Saldo previsto")] +
    [(lab, r_(ent), r_(sai), r_(saldo)) for lab, ent, sai, saldo in proj_ac],
    [10, 18, 18, 18], titulo="Previsão de Fluxo de Caixa — Próximos 6 Meses")

# saldos
aba("Saldo contas",
    [("Conta", f"Saldo em {FIM_MES_ANT_LABEL}")] +
    [(n, r_(v)) for n, v in saldos_conta] +
    [("Total", r_(total_saldos))],
    [30, 18], titulo=f"Saldo nas contas em {FIM_MES_ANT_LABEL}")

wb.save(ARQ_XLSX)
print(f"OK: {ARQ_XLSX}")
