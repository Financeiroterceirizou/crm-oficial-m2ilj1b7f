#!/usr/bin/env python3
# DRE GERENCIAL — UNILAGES (CC SAO JOAQUIM) — v1.1, 2026-10-15
#   v1.1 (pedido Vinícius 15/10): seção DESPESAS DETALHADAS — despesa por despesa agrupada por
#     categoria (data, descrição, conta, valor), no PDF (página própria) e aba no Excel.
#   Formato: entradas e saídas por categoria NA ORDEM DO PLANO DE CONTAS do Controlle
#   (árvore pai→filho, raízes e filhos ordenados pelo CÓDIGO da categoria), só categorias com
#   movimentação no mês de referência. Regime de CAIXA (situation in 1,2, mês por dt_billing).
#   CC: SAO JOAQUIM (170564). Referência inicial: abril/2026 (pedido Vinícius 15/10 — modelo).
# Uso: python3 dre_unilages.py [YYYY-MM-DD-geracao] [YYYY-MM-DD-do-mes-ref]  (default: hoje, mês anterior)
import json, os, sys, urllib.request, time
from collections import defaultdict
from datetime import date, timedelta

from reportlab.lib.pagesizes import A4
from reportlab.lib.units import cm
from reportlab.lib import colors
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.platypus import (SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle,
                                PageBreak, Image as RLImage, KeepTogether)

BASE = "https://api-v1.controlle.com"
_dir = os.path.dirname(os.path.abspath(__file__))
TOKEN = os.environ.get("CONTROLLE_TOKEN_UNILAGES") or open(os.path.join(_dir, ".controlle_token")).read().strip()
UA = {"Authorization": f"Bearer {TOKEN}", "User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) Chrome/126.0"}
LARANJA = colors.HexColor("#ff501c")
LARANJA_CLARO = colors.HexColor("#ffe3d6")
PRETO = colors.HexColor("#1a1a1a")
CINZA = colors.HexColor("#f5f5f5")
VERDE = colors.HexColor("#1a7f37")
VERMELHO = colors.HexColor("#c0392b")
LOGO = os.path.join(_dir, "logo-terceirizou.png")
SIMBOLO = os.path.join(_dir, "simbolo-terceirizou.png")

CC_ID, CC_NOME = 170564, "SAO JOAQUIM"

MES_PT = {1:"janeiro",2:"fevereiro",3:"março",4:"abril",5:"maio",6:"junho",7:"julho",8:"agosto",9:"setembro",10:"outubro",11:"novembro",12:"dezembro"}

def req(url, tries=4):
    for i in range(tries):
        try:
            r = urllib.request.Request(url)
            for k, v in UA.items():
                r.add_header(k, v)
            with urllib.request.urlopen(r, timeout=120) as resp:
                return json.loads(resp.read().decode("utf-8", "replace"))
        except urllib.error.HTTPError as e:
            if e.code in (500, 502, 503) and i < tries - 1:
                time.sleep(3 * (i + 1))
                continue
            raise

def tx_list(start, end, **filtros):
    out, page = [], 1
    while True:
        url = (f"{BASE}/transaction/v1/transactions/list?start_date={start}&end_date={end}"
               f"&page={page}&orderBy=date&orderByCardinality=ASC")
        for k, v in filtros.items():
            url += f"&{k}={v}"
        tl = req(url).get("results", {}).get("transactionsList", [])
        out.extend(tl)
        if len(tl) < 100:
            return out
        page += 1

def brl(cents):
    v = cents / 100
    s = f"{abs(v):,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")
    return ("-" if v < 0 else "") + "R$ " + s

def bdate(t):
    return (t.get("dt_billing") or t.get("dt_due") or "")[:10]

def cat_id_nome(t):
    cats = t.get("apportionments_plan_account") or []
    if not cats:
        return None, "(sem categoria)"
    return cats[0].get("id_category"), (cats[0].get("ds_category") or "?")

# ===== datas =====
HOJE = date.fromisoformat(sys.argv[1]) if len(sys.argv) > 1 else date.today()
MES_REF = date.fromisoformat(sys.argv[2]) if len(sys.argv) > 2 else date(HOJE.year, HOJE.month, 1) - timedelta(days=1)
MES_REF = MES_REF.replace(day=1)
FIM_MES = MES_REF.replace(day=28) + timedelta(days=4)
FIM_MES = FIM_MES - timedelta(days=FIM_MES.day)

HOJE_LABEL = HOJE.strftime("%d/%m/%Y")

# ===== plano de contas (ordem estrutural da árvore) =====
plano = [p for p in req(f"{BASE}/plan-account/v1/planAccountsEntities").get("results", []) if p.get("status") == 1]
idx = {p["id"]: p for p in plano}
by_parent = defaultdict(list)
roots = []
for p in plano:
    pid = p.get("id_plan_accounts_parent")
    if pid is None:
        roots.append(p)
    else:
        by_parent[pid].append(p)

# ordem de exibição da árvore (pré-ordem) — raízes e filhos ordenados pelo CÓDIGO da categoria
# (prefixo numérico do nome, ex. "01", "20.03"); sem código, vai ao fim mantendo a ordem da API
import re
def _cod(p):
    m = re.match(r"^(\d+)", p["ds_category"] or "")
    return (int(m.group(1)),) if m else (9999,)
roots.sort(key=_cod)
for pid in by_parent:
    by_parent[pid].sort(key=_cod)
ordem = []  # (id, nome, level, is_father)
def walk(node):
    ordem.append((node["id"], node["ds_category"], node["level"], node["is_father"]))
    for ch in by_parent.get(node["id"], []):
        walk(ch)
for r in roots:
    walk(r)

# ===== lançamentos do mês no CC =====
_tx = tx_list(MES_REF.isoformat(), FIM_MES.isoformat())
mes_tx = [t for t in _tx if t.get("situation") in (1, 2)
          and any(c.get("cost_center_id") == CC_ID for c in (t.get("apportionments_cost_center") or []))]
# excluir transferências entre contas (categoria 99.01)
mes_tx = [t for t in mes_tx if not any((c.get("ds_category") or "").startswith("99.01")
                                       for c in (t.get("apportionments_plan_account") or []))]

# valores por categoria (id da categoria)
valores = defaultdict(lambda: [0, 0])  # id → [valor, n_lanc]
for t in mes_tx:
    cid, nome = cat_id_nome(t)
    v = t["value_in_cent"]
    valores[cid][0] += v
    valores[cid][1] += 1

# categorias com movimento
com_mov = {cid for cid, (v, n) in valores.items() if v != 0}

# ===== montagem da DRE =====
linhas = []  # (tipo, nome, valor, n) tipo: grupo | cat
for cid, nome, level, is_father in ordem:
    if is_father:
        filhos = by_parent.get(cid, [])
        if any(ch["id"] in com_mov for ch in filhos):
            linhas.append(("grupo", nome, None, None))
    else:
        if cid in com_mov:
            v, n = valores[cid]
            linhas.append(("cat", nome, v, n))
# categorias com movimento que NÃO estão no plano ativo (ex.: categoria cancelada com lançamento no mês)
ids_no_plano = {cid for cid, _, _, _ in ordem}
for cid in sorted(com_mov - ids_no_plano, key=lambda x: valores[x][0]):
    v, n = valores[cid]
    linhas.append(("cat", "(categoria fora do plano ativo)", v, n))

tot_rec = sum(v for _, (v, n) in valores.items() if v > 0)
tot_desp = sum(v for _, (v, n) in valores.items() if v < 0)
resultado = tot_rec + tot_desp

# ===== PDF =====
styles = getSampleStyleSheet()
h1 = ParagraphStyle("h1", parent=styles["Heading1"], fontSize=16, textColor=PRETO, spaceAfter=4)
h2 = ParagraphStyle("h2", parent=styles["Heading2"], fontSize=12, textColor=LARANJA, spaceBefore=14, spaceAfter=3)
h3 = ParagraphStyle("h3", parent=styles["Heading3"], fontSize=10, textColor=PRETO, spaceBefore=10, spaceAfter=3)
sub = ParagraphStyle("sub", parent=styles["Normal"], fontName="Helvetica-Oblique", fontSize=8.5, textColor=colors.HexColor("#777777"), spaceAfter=10)
cell = ParagraphStyle("cell", parent=styles["Normal"], fontSize=8)
cellb = ParagraphStyle("cellb", parent=styles["Normal"], fontSize=8, fontName="Helvetica-Bold")
cellr = ParagraphStyle("cellr", parent=cell, alignment=2)
cellrb = ParagraphStyle("cellrb", parent=cellb, alignment=2)
cellc = ParagraphStyle("cellc", parent=cell, alignment=1)
P = Paragraph

def P_val(cents, style, bold=False):
    st = ParagraphStyle("v", parent=style, textColor=(VERDE if cents > 0 else VERMELHO if cents < 0 else PRETO),
                        fontName="Helvetica-Bold" if bold else "Helvetica")
    return P(brl(cents), st)

def tabela(rows, widths, fs=8):
    t = Table(rows, colWidths=widths, repeatRows=1)
    style = [("FONTNAME", (0,0), (-1,-1), "Helvetica"), ("FONTSIZE", (0,0), (-1,-1), fs),
             ("TOPPADDING", (0,0), (-1,-1), 3), ("BOTTOMPADDING", (0,0), (-1,-1), 3),
             ("VALIGN", (0,0), (-1,-1), "MIDDLE"),
             ("BACKGROUND", (0,0), (-1,0), LARANJA), ("TEXTCOLOR", (0,0), (-1,0), colors.white),
             ("FONTNAME", (0,0), (-1,0), "Helvetica-Bold"),
             ("LINEBELOW", (0,-1), (-1,-1), 0.7, LARANJA)]
    for i in range(1, len(rows)):
        if i % 2 == 0:
            style.append(("BACKGROUND", (0,i), (-1,i), CINZA))
    t.setStyle(TableStyle(style))
    return t

def rodape_simbolo(canvas, doc):
    if os.path.exists(SIMBOLO):
        alt = 1.2 * cm
        canvas.drawImage(SIMBOLO, 18.4*cm, 1.1*cm, width=alt*372/553, height=alt, mask="auto")

ARQ_PDF = f"artifacts/{HOJE.strftime('%y%m%d')}_DRE_Unilages_Sao_Joaquim_{MES_REF.strftime('%Y%m')}.pdf"
doc = SimpleDocTemplate(ARQ_PDF, pagesize=A4, leftMargin=1.5*cm, rightMargin=1.5*cm,
                        topMargin=1.3*cm, bottomMargin=1.6*cm,
                        title=f"DRE Gerencial — Unilages {CC_NOME} — {MES_PT[MES_REF.month].capitalize()} de {MES_REF.year}")
E = []
h1c = ParagraphStyle("h1c", parent=h1, fontSize=13, alignment=1, spaceAfter=1)
h2c = ParagraphStyle("h2c", parent=h2, fontSize=11, alignment=1, spaceBefore=2, spaceAfter=2)
subc = ParagraphStyle("subc", parent=sub, fontSize=8, alignment=1, spaceAfter=0)
if os.path.exists(LOGO):
    img = RLImage(LOGO, width=5.5*cm, height=5.5*cm*561/1600)
    img.hAlign = "CENTER"
    E.append(img)
E.append(Spacer(1, 8))
E.append(P("<b>UNILAGES — DRE GERENCIAL</b>", h1c))
E.append(P(f"Centro de Custo {CC_NOME}", h2c))
E.append(P(f"Gerado em {HOJE_LABEL} · Fonte: Controlle · Ref.: {MES_PT[MES_REF.month].capitalize()} de {MES_REF.year} — regime de caixa · Categorias na ordem do plano de contas", subc))
E.append(Spacer(1, 6))

dre_rows = [[P("<b>Categoria</b>", cell), P("<b>Lançamentos</b>", cellc), P("<b>Valor</b>", cellr)]]
for tipo, nome, v, n in linhas:
    if tipo == "grupo":
        dre_rows.append([P(f"<b>{nome}</b>", cellb), P("", cellc), P("", cellr)])
    else:
        dre_rows.append([P("      " + nome, cell), P(str(n), cellc), P_val(v, cellr)])
dre_rows.append([P("<b>Total de Receitas</b>", cellrb), P(f"<b>{sum(n for _, (v, n) in valores.items() if v > 0)}</b>", cellc), P_val(tot_rec, cellrb, True)])
dre_rows.append([P("<b>Total de Despesas</b>", cellrb), P(f"<b>{sum(n for _, (v, n) in valores.items() if v < 0)}</b>", cellc), P_val(tot_desp, cellrb, True)])
dre_rows.append([P("<b>Resultado do mês</b>", cellrb), P(f"<b>{len(mes_tx)}</b>", cellc), P_val(resultado, cellrb, True)])
t = tabela(dre_rows, [11*cm, 2.5*cm, 3*cm])
t.setStyle(TableStyle([("BACKGROUND", (0,len(dre_rows)-3), (-1,len(dre_rows)-1), LARANJA_CLARO)]))
E.append(t)

# ===== Despesas detalhadas (despesa por despesa) =====
E.append(PageBreak())
E.append(P(f"Despesas Detalhadas — {MES_PT[MES_REF.month].capitalize()} de {MES_REF.year} (regime de caixa)", h2))
desp_tx = sorted([t for t in mes_tx if t["value_in_cent"] < 0], key=bdate)
por_cat_d = defaultdict(list)
for t in desp_tx:
    _, nome = cat_id_nome(t)
    por_cat_d[nome].append(t)
for cat in sorted(por_cat_d):
    txs_c = sorted(por_cat_d[cat], key=bdate)
    rows = [[P("<b>Data</b>", cell), P("<b>Descrição</b>", cell), P("<b>Conta</b>", cell), P("<b>Valor</b>", cellr)]]
    for t in txs_c:
        d = bdate(t)
        rows.append([P(f"{d[8:10]}/{d[5:7]}/{d[:4]}", cell), P((t.get("ds_transaction") or "")[:70], cell),
                     P(t.get("ds_account_main") or "", cell), P_val(t["value_in_cent"], cellr)])
    rows.append([P("<b>Total</b>", cellrb), P(f"<b>{cat}</b>", cellrb), P(f"<b>{len(txs_c)}</b>", cellc), P_val(sum(t["value_in_cent"] for t in txs_c), cellrb)])
    tt = tabela(rows, [2.2*cm, 8.3*cm, 3.5*cm, 3*cm])
    tt.setStyle(TableStyle([("BACKGROUND", (0,len(rows)-1), (-1,len(rows)-1), LARANJA_CLARO)]))
    E.append(KeepTogether([P(f"<b>{cat}</b>", h3), tt]))
tot_d = sum(t["value_in_cent"] for t in desp_tx)
td_rows = [[P("<b>Total de Despesas do mês</b>", cellrb), P(f"<b>{len(desp_tx)}</b>", cellc), P_val(tot_d, cellrb, True)]]
t = tabela(td_rows, [11*cm, 2.5*cm, 3*cm])
t.setStyle(TableStyle([("BACKGROUND", (0,0), (-1,0), LARANJA_CLARO)]))
E.append(t)

E.append(Spacer(1, 10))
E.append(P("Gerado automaticamente pela Terceirizou · dados do Controlle", sub))
doc.build(E, onFirstPage=rodape_simbolo, onLaterPages=rodape_simbolo)
print(f"OK: {ARQ_PDF}")

# ===== Excel =====
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment

ARQ_XLSX = f"artifacts/{HOJE.strftime('%y%m%d')}_DRE_Unilages_Sao_Joaquim_{MES_REF.strftime('%Y%m')}.xlsx"
wb = Workbook()
ws = wb.active
ws.title = ("DRE " + MES_REF.strftime("%m%y"))[:31]
ws.append((f"DRE Gerencial — Unilages {CC_NOME} — {MES_PT[MES_REF.month].capitalize()} de {MES_REF.year} (regime de caixa)",))
ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=3)
c = ws.cell(row=1, column=1)
c.font = Font(bold=True, size=12, color="FF501C")
c.alignment = Alignment(horizontal="center", vertical="center")
ws.row_dimensions[1].height = 22
ws.append(("Categoria", "Lançamentos", "Valor"))
FILL_H = PatternFill("solid", fgColor="FF501C")
FILL_T = PatternFill("solid", fgColor="FFE3D6")
FH = Font(bold=True, color="FFFFFF")
FB = Font(bold=True)
VERDE_XL = "1A7F37"
VERMELHO_XL = "C0392B"
for cc in ws[2]:
    cc.fill = FILL_H
    cc.font = FH
    cc.alignment = Alignment(horizontal="center", vertical="center")
for tipo, nome, v, n in linhas:
    if tipo == "grupo":
        ws.append((nome, "", ""))
        r = ws.max_row
        for cc in ws[r]:
            cc.font = Font(bold=True)
            cc.fill = FILL_T
    else:
        ws.append(("    " + nome, n, v / 100))
ws.append(("Total de Receitas", sum(n for _, (v, n) in valores.items() if v > 0), tot_rec / 100))
ws.append(("Total de Despesas", sum(n for _, (v, n) in valores.items() if v < 0), tot_desp / 100))
ws.append(("Resultado do mês", len(mes_tx), resultado / 100))
for r in range(ws.max_row - 2, ws.max_row + 1):
    for cc in ws[r]:
        cc.font = Font(bold=True)
        cc.fill = FILL_T
for row in ws.iter_rows(min_row=3):
    cc = row[2]
    if isinstance(cc.value, (int, float)):
        cc.number_format = '"R$" #,##0.00'
        cc.font = Font(bold=bool(row[0].font and row[0].font.bold),
                       color=(VERDE_XL if cc.value > 0 else VERMELHO_XL if cc.value < 0 else "1A1A1A"))
    if isinstance(row[1].value, (int, float)):
        row[1].number_format = "0"
ws.column_dimensions["A"].width = 50
ws.column_dimensions["B"].width = 14
ws.column_dimensions["C"].width = 16
ws.freeze_panes = "A3"

# aba de despesas detalhadas (despesa por despesa)
ws2 = wb.create_sheet(("Detalhe Despesas " + MES_REF.strftime("%m%y"))[:31])
ws2.append((f"Despesas Detalhadas — Unilages {CC_NOME} — {MES_PT[MES_REF.month].capitalize()} de {MES_REF.year} (regime de caixa)",))
ws2.merge_cells(start_row=1, start_column=1, end_row=1, end_column=5)
c2 = ws2.cell(row=1, column=1)
c2.font = Font(bold=True, size=12, color="FF501C")
c2.alignment = Alignment(horizontal="center", vertical="center")
ws2.row_dimensions[1].height = 22
ws2.append(("Categoria", "Data", "Descrição", "Conta", "Valor"))
for cc in ws2[2]:
    cc.fill = FILL_H
    cc.font = FH
    cc.alignment = Alignment(horizontal="center", vertical="center")
for cat in sorted(por_cat_d):
    for t in sorted(por_cat_d[cat], key=bdate):
        d = bdate(t)
        ws2.append((cat, f"{d[8:10]}/{d[5:7]}/{d[:4]}", (t.get("ds_transaction") or "")[:80],
                    t.get("ds_account_main") or "", t["value_in_cent"] / 100))
    ws2.append((f"Total {cat}", "", "", "", sum(t["value_in_cent"] for t in por_cat_d[cat]) / 100))
    r = ws2.max_row
    for cc in ws2[r]:
        cc.font = Font(bold=True)
        cc.fill = FILL_T
ws2.append(("Total de Despesas do mês", "", "", "", tot_d / 100))
r = ws2.max_row
for cc in ws2[r]:
    cc.font = Font(bold=True)
    cc.fill = FILL_T
for row in ws2.iter_rows(min_row=3):
    cc = row[4]
    if isinstance(cc.value, (int, float)):
        cc.number_format = '"R$" #,##0.00'
        cc.font = Font(bold=bool(row[0].font and row[0].font.bold),
                       color=(VERDE_XL if cc.value > 0 else VERMELHO_XL if cc.value < 0 else "1A1A1A"))
for col, w in zip("ABCDE", (40, 12, 55, 24, 16)):
    ws2.column_dimensions[col].width = w
ws2.freeze_panes = "A3"
wb.save(ARQ_XLSX)
print(f"OK: {ARQ_XLSX}")
