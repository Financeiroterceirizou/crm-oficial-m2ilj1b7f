#!/usr/bin/env python3
"""Consolida os pacotes DRE Unilages (mai/jun/jul 2026) em 1 Excel.
Abas: Resumo + para cada CC (GERAL, Correia Pinto, Curitibanos, Lages+BM Adm,
Otacilio Costa, Sao Joaquim, Urubici) x 3 meses: DRE + Detalhe Despesas."""
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from copy import copy
import os

PACOTES = [
    ("Geral", "GERAL"),
    ("correia_pinto", "CORREIA PINTO"),
    ("curitibanos", "CURITIBANOS"),
    ("lages_bm_adm", "LAGES + BM ADM"),
    ("otacilio_costa", "OTACILIO COSTA"),
    ("sao_joaquim", "SAO JOAQUIM"),
    ("urubici", "URUBICI"),
]
MESES = [("202605", "05-26"), ("202606", "06-26"), ("202607", "07-26")]
ARQ = "artifacts/261022_DRE_Unilages_Completo_Mai_Jun_Jul_2026.xlsx"

wb = Workbook()
wb.remove(wb.active)

# ---- Resumo ----
ws = wb.create_sheet("Resumo")
ws.append(("Resumo — Resultado do mês por centro de custo (regime de caixa, fonte Controlle)",))
ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=7)
c = ws.cell(row=1, column=1)
c.font = Font(bold=True, size=12, color="FF501C")
c.alignment = Alignment(horizontal="center", vertical="center")
ws.row_dimensions[1].height = 22
ws.append(("Centro de custo", "Mês", "Receitas", "Despesas", "Resultado", "Lançamentos", ""))
for cc in ws[2]:
    cc.fill = PatternFill("solid", fgColor="FF501C")
    cc.font = Font(bold=True, color="FFFFFF")
    cc.alignment = Alignment(horizontal="center", vertical="center")

resumo = []  # (cc_label, mes, rec, desp, res, n)
for tag, label in PACOTES:
    for m, mm in MESES:
        src = load_workbook(f"artifacts/261022_DRE_Unilages_{tag}_{m}.xlsx")
        dre = src.worksheets[0]
        vals = {}
        n = None
        for row in dre.iter_rows(min_row=3, max_col=3):
            a = row[0].value
            if a in ("Total de Receitas", "Total de Despesas", "Resultado do mês"):
                vals[a] = row[2].value
            if a == "Resultado do mês":
                n = row[1].value
        resumo.append((label, mm, vals.get("Total de Receitas"), vals.get("Total de Despesas"),
                       vals.get("Resultado do mês"), n))
for label, mm, rec, desp, res, n in resumo:
    ws.append((label, mm, rec, desp, res, n, ""))
    r = ws.max_row
    for col in (3, 4, 5):
        cc = ws.cell(row=r, column=col)
        cc.number_format = '"R$" #,##0.00'
        if isinstance(cc.value, (int, float)):
            cc.font = Font(color=("1A7F37" if col == 5 and cc.value > 0 else "C0392B" if col == 5 and cc.value < 0 else "1A1A1A"))
    ws.cell(row=r, column=6).number_format = "0"
for col, w in zip("ABCDEFG", (22, 9, 16, 16, 16, 14, 2)):
    ws.column_dimensions[col].width = w
ws.freeze_panes = "A3"

# ---- Abas DRE + Detalhe por pacote/mês ----
def copia_sheet(orig, dest, nome):
    d = wb.create_sheet(nome)
    for row in orig.iter_rows():
        for cell in row:
            nc = d.cell(row=cell.row, column=cell.column, value=cell.value)
            if cell.has_style:
                nc.font = copy(cell.font)
                nc.fill = copy(cell.fill)
                nc.number_format = cell.number_format
                nc.alignment = copy(cell.alignment)
                nc.border = copy(cell.border)
    for k, dim in orig.column_dimensions.items():
        d.column_dimensions[k].width = dim.width
    for k, dim in orig.row_dimensions.items():
        d.row_dimensions[k].height = dim.height
    for rng in orig.merged_cells.ranges:
        d.merge_cells(str(rng))
    if orig.freeze_panes:
        d.freeze_panes = orig.freeze_panes
    return d

for tag, label in PACOTES:
    for m, mm in MESES:
        src = load_workbook(f"artifacts/261022_DRE_Unilages_{tag}_{m}.xlsx")
        copia_sheet(src.worksheets[0], None, f"DRE {mm} {label}"[:31])
        if len(src.worksheets) > 1:
            copia_sheet(src.worksheets[1], None, f"Det {mm} {label}"[:31])

wb.save(ARQ)
print(f"OK: {ARQ} — {len(wb.sheetnames)} abas")