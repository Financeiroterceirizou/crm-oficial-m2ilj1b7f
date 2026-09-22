#!/usr/bin/env python3
# Relatório Gerencial Mensal — TERCEIRIZOU (v4.2, 2026-09-18)
# Uso: python3 relatorio_mensal_terceirizou_v4.py [YYYY-MM-DD]
#   Sem argumento: mês anterior completo (rodar no dia 05 via cron).
#   Com argumento: último dia do mês de referência (ex.: 2026-08-31).
# Fonte: API Controlle v1. Token: scripts/.controlle_token ou env CONTROLLE_TOKEN.
# Gera PDF (logo + laranja #ff501c) + Excel (7 abas).
# Seções: Resumo · Fluxo de Caixa realizado (planilha) · Categorias Receitas e Despesas
# (planilha única com resultado) · DRE caixa por grupo · DRE competência no mesmo formato
# do caixa (janela larga jan/ano→dez/ano+1 + filtro dt_competence) · Receitas em aberto por
# cliente · Saldo nas contas · Projeção dos PRÓXIMOS 12 meses (balancePreview) com gráfico.
import json, os, sys, urllib.request
from collections import defaultdict
from datetime import date, timedelta

from reportlab.lib.pagesizes import A4
from reportlab.lib.units import cm
from reportlab.lib import colors
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.platypus import (SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle,
                                Image)
from reportlab.graphics.shapes import Drawing, String
from reportlab.graphics.charts.barcharts import VerticalBarChart

BASE = "https://api-v1.controlle.com"
_token_path = os.path.join(os.path.dirname(os.path.abspath(__file__)), ".controlle_token")
TOKEN = os.environ.get("CONTROLLE_TOKEN") or (open(_token_path).read().strip() if os.path.exists(_token_path) else "")
UA = {"Authorization": f"Bearer {TOKEN}", "User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) Chrome/126.0"}
LARANJA = colors.HexColor("#ff501c")
LARANJA_CLARO = colors.HexColor("#ffe3d6")
PRETO = colors.HexColor("#1a1a1a")
CINZA = colors.HexColor("#f5f5f5")

MES_PT = {1:"janeiro",2:"fevereiro",3:"março",4:"abril",5:"maio",6:"junho",7:"julho",8:"agosto",9:"setembro",10:"outubro",11:"novembro",12:"dezembro"}
MES_AB = {1:"jan",2:"fev",3:"mar",4:"abr",5:"mai",6:"jun",7:"jul",8:"ago",9:"set",10:"out",11:"nov",12:"dez"}

def req(url):
    r = urllib.request.Request(url)
    for k, v in UA.items():
        r.add_header(k, v)
    with urllib.request.urlopen(r, timeout=90) as resp:
        return json.loads(resp.read().decode("utf-8", "replace"))

def tx_list(start, end, **filtros):
    out, page = [], 1
    while True:
        url = (f"{BASE}/transaction/v1/transactions/list?start_date={start}&end_date={end}"
               f"&page={page}&orderBy=date&orderByCardinality=ASC")
        for k, v in filtros.items():
            url += f"&{k}={v}"
        tl = req(url).get("results", {}).get("transactionsList", [])
        out.extend(tl)
        if len(tl) < 100:
            return out
        page += 1

def brl(cents):
    v = cents / 100
    s = f"{abs(v):,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")
    return ("-" if v < 0 else "") + "R$ " + s

def add_months(d, n):
    m = d.month - 1 + n
    y = d.year + m // 12
    m = m % 12 + 1
    return date(y, m, 1)

# ===== período de referência =====
if len(sys.argv) > 1:
    MES_FIM = date.fromisoformat(sys.argv[1])
else:
    MES_FIM = date.today().replace(day=1) - timedelta(days=1)
fim = MES_FIM.isoformat()
MES_INI = MES_FIM.replace(day=1).isoformat()
MES_LABEL = f"{MES_PT[MES_FIM.month].capitalize()} de {MES_FIM.year}"
FIM_LABEL = MES_FIM.strftime("%d/%m/%Y")

# ===== dados =====
bal = req(f"{BASE}/transaction/v1/transactions/balances?start_date={MES_INI}&end_date={fim}")["results"]
entradas_mes, saidas_mes = bal["revenuesDone"], bal["expensesDone"]
saldo_anterior, saldo_final = bal["previousMonthBalance"], bal["balanceDone"]
resultado_mes = entradas_mes + saidas_mes

# mapa categoria -> grupo pai (para o DRE caixa)
plan = req(f"{BASE}/plan-account/v1/planAccountsEntities").get("results", [])
pai_nome, cat_grupo = {}, {}
for c in plan:
    pai_nome[c["id"]] = c["ds_category"]
for c in plan:
    if c.get("id_plan_accounts_parent"):
        cat_grupo[c["id"]] = pai_nome.get(c["id_plan_accounts_parent"], "?")

txs = tx_list(MES_INI, fim)
normais = [t for t in txs if not (t.get("ds_transaction") or "").startswith("Transferência")]
rec_vals, desp_cat = defaultdict(int), defaultdict(int)
grupo_val = defaultdict(int)
for t in normais:
    for c in (t.get("apportionments_plan_account") or []):
        v = c.get("value") or 0
        nome = c.get("ds_category") or "?"
        if v > 0:
            rec_vals[nome] += v
        elif v < 0:
            desp_cat[nome] += v
        grupo_val[cat_grupo.get(c.get("id_category"), nome)] += v

# DRE competência: TODAS as transações com competência no mês (pagas + não pagas), por categoria.
# IMPORTANTE: a API filtra pela data de vencimento — lançamentos com competência no mês mas
# vencimento distante (ex.: passagens pagas meses depois) só aparecem numa JANELA LARGA
# (ano de referência até ano seguinte). Validado contra o PDF do sistema (18/09): ago/26
# entradas 67.754,02 / saídas -82.873,38 / resultado -15.119,36 — match exato.
comp_rec, comp_desp = defaultdict(int), defaultdict(int)
cat_grupo_id = {}  # nome da categoria -> grupo pai (para o DRE competência por grupo)
for t in tx_list(f"{MES_FIM.year}-01-01", f"{MES_FIM.year + 1}-12-31"):
    if t["dt_competence"][:7] != MES_FIM.strftime("%Y-%m"):
        continue
    if (t.get("ds_transaction") or "").startswith("Transferência"):
        continue
    for c in (t.get("apportionments_plan_account") or []):
        v = c.get("value") or 0
        nome = c.get("ds_category") or "?"
        gid = cat_grupo.get(c.get("id_category"), nome)
        cat_grupo_id[nome] = gid
        if v > 0:
            comp_rec[nome] += v
        elif v < 0:
            comp_desp[nome] += v
comp_te, comp_ts = sum(comp_rec.values()), sum(comp_desp.values())

# receitas em aberto por cliente (desde 2017)
abertas = tx_list("2017-01-01", fim, **{"activity_type": "1", "situation": "[0]"})
por_cliente = defaultdict(int)
for t in abertas:
    nome = (t.get("name_contact") or "").strip() or (t.get("ds_transaction") or "").strip()[:45]
    if nome.startswith("Transferência"):
        continue
    por_cliente[nome] += t.get("value_in_cent") or 0
total_aberto = sum(por_cliente.values())

# saldos por conta no fim do mês
contas = req(f"{BASE}/account/v1/accounts").get("results", [])
saldos_conta = []
for c in contas:
    if c.get("status") != 1:
        continue
    b = req(f"{BASE}/transaction/v1/transactions/balances?start_date=2017-01-01&end_date={fim}&id_account_main={c['id']}")["results"]
    if b["balanceDone"] != 0:
        saldos_conta.append((c["ds_account"], b["balanceDone"]))

# ===== PDF =====
styles = getSampleStyleSheet()
h1 = ParagraphStyle("h1", parent=styles["Heading1"], fontSize=15, textColor=PRETO, spaceAfter=2)
h2 = ParagraphStyle("h2", parent=styles["Heading2"], fontSize=11.5, textColor=LARANJA, spaceBefore=12, spaceAfter=5)
sub = ParagraphStyle("sub", parent=styles["Normal"], fontSize=9, textColor=colors.HexColor("#666666"), spaceAfter=8)
body = ParagraphStyle("body", parent=styles["Normal"], fontSize=9, leading=12.5)
cell = ParagraphStyle("cell", parent=styles["Normal"], fontSize=8)
cellb = ParagraphStyle("cellb", parent=styles["Normal"], fontSize=8, fontName="Helvetica-Bold")
cellr = ParagraphStyle("cellr", parent=cell, alignment=2)
cellrb = ParagraphStyle("cellrb", parent=cellb, alignment=2)

def tabela(rows, widths):
    t = Table(rows, colWidths=widths, repeatRows=1)
    style = [("FONTNAME", (0,0), (-1,-1), "Helvetica"), ("FONTSIZE", (0,0), (-1,-1), 8),
             ("TOPPADDING", (0,0), (-1,-1), 4), ("BOTTOMPADDING", (0,0), (-1,-1), 4),
             ("VALIGN", (0,0), (-1,-1), "MIDDLE"),
             ("BACKGROUND", (0,0), (-1,0), LARANJA), ("TEXTCOLOR", (0,0), (-1,0), colors.white),
             ("FONTNAME", (0,0), (-1,0), "Helvetica-Bold"),
             ("LINEBELOW", (0,-1), (-1,-1), 0.7, LARANJA)]
    for i in range(1, len(rows)):
        if i % 2 == 0:
            style.append(("BACKGROUND", (0,i), (-1,i), CINZA))
    t.setStyle(TableStyle(style))
    return t

P = Paragraph
ARQ = f"artifacts/relatorio-terceirizou-{MES_FIM.strftime('%Y-%m')}.pdf"
doc = SimpleDocTemplate(ARQ, pagesize=A4, leftMargin=1.6*cm, rightMargin=1.6*cm, topMargin=1.4*cm, bottomMargin=1.4*cm)
E = []

logo_path = os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", "artifacts", "logo-terceirizou.png")
if os.path.exists(logo_path):
    E.append(Image(logo_path, width=6.2*cm, height=6.2*cm*702/2000))
E.append(Spacer(1, 4))
E.append(P("<b>Relatório Gerencial Mensal</b>", h1))
E.append(P(f"{MES_LABEL} · Fechamento · Fonte: Controlle · Gerado em {date.today().strftime('%d/%m/%Y')}", sub))

# 1. Resumo do mês
E.append(P("Resumo do mês", h2))
resumo = [
    [P("<b>Indicador</b>", cell), P(f"<b>{MES_LABEL}</b>", cellr)],
    [P("Entradas (realizado)", cell), P(brl(entradas_mes), cellr)],
    [P("Saídas (realizado)", cell), P(brl(saidas_mes), cellr)],
    [P("<b>Resultado</b>", cellrb), P(f"<b>{brl(resultado_mes)}</b>", cellrb)],
    [P(f"Saldo em {FIM_LABEL}", cell), P(brl(saldo_final), cellr)],
]
E.append(tabela(resumo, [7*cm, 5*cm]))

# 2. Fluxo de Caixa — realizado (planilha com resultado)
E.append(P("Fluxo de Caixa — realizado", h2))
fc_rows = [[P("<b>Descrição</b>", cell), P(f"<b>{MES_LABEL}</b>", cellr)]]
fc_rows.append([P("Saldo anterior", cell), P(brl(saldo_anterior), cellr)])
fc_rows.append([P("(+) Total de entradas", cell), P(brl(entradas_mes), cellr)])
fc_rows.append([P("(-) Total de saídas", cell), P(brl(saidas_mes), cellr)])
fc_rows.append([P("<b>Resultado do mês</b>", cellrb), P(f"<b>{brl(resultado_mes)}</b>", cellrb)])
fc_rows.append([P("<b>Saldo final</b>", cellrb), P(f"<b>{brl(saldo_final)}</b>", cellrb)])
t = tabela(fc_rows, [7*cm, 5*cm])
t.setStyle(TableStyle([("BACKGROUND", (0,4), (-1,5), LARANJA_CLARO)]))
E.append(t)

# 3. Categorias Receitas e Despesas (planilha única com resultado)
E.append(P("Categorias Receitas e Despesas", h2))
cd_rows = [[P("<b>Categoria</b>", cell), P("<b>Entradas</b>", cellr), P("<b>Saídas</b>", cellr)]]
for nome, v in sorted(rec_vals.items(), key=lambda x: -x[1]):
    cd_rows.append([P(nome, cell), P(brl(v), cellr), P("—", cellr)])
for nome, v in sorted(desp_cat.items(), key=lambda x: x[1]):
    cd_rows.append([P(nome, cell), P("—", cellr), P(brl(v), cellr)])
cd_rows.append([P("<b>Totais</b>", cellrb), P(f"<b>{brl(entradas_mes)}</b>", cellrb), P(f"<b>{brl(saidas_mes)}</b>", cellrb)])
cd_rows.append([P("", cell), P(""), P(f"<b>{brl(resultado_mes)}</b>", cellrb)])
t = tabela(cd_rows, [9*cm, 3.5*cm, 3.5*cm])
t.setStyle(TableStyle([("BACKGROUND", (0,len(cd_rows)-2), (-1,len(cd_rows)-1), LARANJA_CLARO)]))
E.append(t)

# 4. DRE Gerencial — regime caixa (por grupo)
E.append(P("DRE Gerencial — regime caixa", h2))
dre_rows = [[P("<b>Conta</b>", cell), P(f"<b>{MES_LABEL}</b>", cellr), P("<b>% receita</b>", cellr)]]
dre_rows.append([P("Receita total", cell), P(brl(entradas_mes), cellr), P("100,0%", cellr)])
GRUPOS_ORDEM = ["CUSTOS OPERACIONAIS", "DESPESAS DE RH", "DESPESAS ADMINISTRATIVAS E COMERCIAS",
                "IMPOSTOS SOBRE FATURAMENTO", "DESPESAS FINANCEIRAS"]
GRUPO_LABEL = {
    "CUSTOS OPERACIONAIS": "Custos Operacionais",
    "DESPESAS DE RH": "Despesas de RH",
    "DESPESAS ADMINISTRATIVAS E COMERCIAS": "Despesas Administrativas e Comerciais",
    "IMPOSTOS SOBRE FATURAMENTO": "Impostos sobre Faturamento",
    "DESPESAS FINANCEIRAS": "Despesas Financeiras",
}
for g in GRUPOS_ORDEM:
    v = grupo_val.get(g, 0)
    if v:
        dre_rows.append([P(f"(-) {GRUPO_LABEL[g]}", cell), P(brl(v), cellr),
                         P(f"{abs(v)/entradas_mes*100:.1f}%".replace(".", ","), cellr)])
dre_rows.append([P("<b>Resultado do mês</b>", cellrb), P(f"<b>{brl(resultado_mes)}</b>", cellrb),
                 P(f"<b>{resultado_mes/entradas_mes*100:.1f}%</b>".replace(".", ","), cellrb)])
t = tabela(dre_rows, [9*cm, 3.5*cm, 3.5*cm])
t.setStyle(TableStyle([("BACKGROUND", (0,len(dre_rows)-1), (-1,len(dre_rows)-1), LARANJA_CLARO)]))
E.append(t)

# 5. DRE Gerencial — regime competência (mesmo formato do caixa, dados de competência)
E.append(P("DRE Gerencial — regime competência (inclui não pagos e não recebidos)", h2))
comp_total = comp_te + comp_ts
dc_rows = [[P("<b>Conta</b>", cell), P(f"<b>{MES_LABEL}</b>", cellr), P("<b>% receita</b>", cellr)]]
dc_rows.append([P("Receita total", cell), P(brl(comp_te), cellr), P("100,0%", cellr)])
GRUPOS_COMP = ["CUSTOS OPERACIONAIS", "DESPESAS DE RH", "DESPESAS ADMINISTRATIVAS E COMERCIAS",
               "IMPOSTOS SOBRE FATURAMENTO", "DESPESAS FINANCEIRAS"]
for g in GRUPOS_COMP:
    v = sum(val for cat, val in comp_desp.items()
            if cat_grupo_id.get(cat) == g or cat == g)
    if v:
        dc_rows.append([P(f"(-) {GRUPO_LABEL[g]}", cell), P(brl(v), cellr),
                        P(f"{abs(v)/comp_te*100:.1f}%".replace(".", ","), cellr)])
dc_rows.append([P("<b>Resultado do mês (competência)</b>", cellrb), P(f"<b>{brl(comp_total)}</b>", cellrb),
                P(f"<b>{comp_total/comp_te*100:.1f}%</b>".replace(".", ","), cellrb)])
t = tabela(dc_rows, [9*cm, 3.5*cm, 3.5*cm])
t.setStyle(TableStyle([("BACKGROUND", (0,len(dc_rows)-1), (-1,len(dc_rows)-1), LARANJA_CLARO)]))
E.append(t)

# 6. Receitas em aberto por cliente
E.append(P(f"Receitas em aberto até {FIM_LABEL}", h2))
E.append(P(f"Total em aberto: <b>{brl(total_aberto)}</b> em {len(por_cliente)} clientes (entradas não pagas desde o início das atividades):", body))
ab_rows = [[P("<b>Cliente</b>", cell), P("<b>Valor em aberto</b>", cellr)]]
for nome, v in sorted(por_cliente.items(), key=lambda x: -x[1]):
    ab_rows.append([P(nome, cell), P(brl(v), cellr)])
ab_rows.append([P("<b>Total</b>", cellrb), P(f"<b>{brl(total_aberto)}</b>", cellrb)])
E.append(tabela(ab_rows, [11*cm, 5*cm]))
E.append(Spacer(1, 4))
E.append(P("Nota: parte desses valores pode ser receita já recebida e não conciliada — validar antes de cobrança.", body))

# 7. Saldo nas contas
E.append(P(f"Saldo nas Contas dia {FIM_LABEL}", h2))
sc_rows = [[P("<b>Conta</b>", cell), P(f"<b>Saldo em {FIM_LABEL}</b>", cellr)]]
for nome, v in saldos_conta:
    sc_rows.append([P(nome, cell), P(brl(v), cellr)])
sc_rows.append([P("<b>Total</b>", cellrb), P(f"<b>{brl(sum(v for _, v in saldos_conta))}</b>", cellrb)])
E.append(tabela(sc_rows, [11*cm, 5*cm]))

# 8. Projeção 12 meses com gráfico — PRÓXIMOS 12 meses (mês seguinte ao de referência)
E.append(P("Projeção de fluxo de caixa — 12 meses", h2))
E.append(P(f"Projeção dos próximos 12 meses (lançamentos previstos do Controlle, base {MES_LABEL}):", body))
proj = []
for i in range(1, 13):
    ini_m = add_months(MES_FIM.replace(day=1), i)
    fim_m = (add_months(ini_m, 1) - timedelta(days=1))
    label = f"{MES_AB[ini_m.month]}/{str(ini_m.year)[2:]}"
    b = req(f"{BASE}/transaction/v1/transactions/balances?start_date={ini_m.isoformat()}&end_date={fim_m.isoformat()}")["results"]
    proj.append((label, b["balancePreview"]))
d = Drawing(17*cm, 5.2*cm)
chart = VerticalBarChart()
chart.x, chart.y, chart.width, chart.height = 42, 14, 430, 120
chart.data = [[v for _, v in proj]]
chart.categoryAxis.categoryNames = [m for m, _ in proj]
chart.categoryAxis.labels.fontName = "Helvetica"
chart.categoryAxis.labels.fontSize = 5.5
chart.categoryAxis.labels.angle = 45
chart.valueAxis.valueMin = 0
chart.valueAxis.valueMax = int(max(5000000, max(v for _, v in proj) * 1.15) / 500000) * 500000
chart.valueAxis.valueStep = 500000
chart.valueAxis.labels.fontName = "Helvetica"
chart.valueAxis.labels.fontSize = 6
chart.valueAxis.labelTextFormat = lambda v: f"{v/100000:.0f}k"  # v em centavos -> milhares
chart.bars[0].fillColor = LARANJA
chart.bars[0].strokeColor = None
d.add(chart)
for i, (m, v) in enumerate(proj):
    d.add(String(48 + i*36.2, 138, f"{v/100:,.0f}".replace(",", ".")[:-1] + "k", fontSize=5.5, fillColor=colors.HexColor("#555555")))
E.append(d)

E.append(Spacer(1, 10))
E.append(P("Gerado automaticamente pela Terceirizou · dados do Controlle", sub))
doc.build(E)

# ===== Excel (mesmos dados, uma aba por seção) =====
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

ARQ_XLSX = f"artifacts/relatorio-terceirizou-{MES_FIM.strftime('%Y-%m')}.xlsx"
wb = Workbook()
wb.remove(wb.active)
FILL_H = PatternFill("solid", fgColor="FF501C")
FILL_T = PatternFill("solid", fgColor="FFE3D6")
FH = Font(bold=True, color="FFFFFF")
FB = Font(bold=True)
TOTALS = ("Total", "Totais", "Resultado", "Saldo final", "Resultado do mês", "Resultado do mês (competência)")

def aba(nome, linhas, larguras):
    ws = wb.create_sheet(nome)
    for r in linhas:
        ws.append(list(r))
    for c in ws[1]:
        c.fill = FILL_H
        c.font = FH
    for row in ws.iter_rows(min_row=2):
        label = str(row[0].value or "")
        if label.startswith(TOTALS):
            for c in row:
                c.font = FB
                c.fill = FILL_T
        for c in row[1:]:
            if isinstance(c.value, (int, float)):
                c.number_format = '"R$" #,##0.00'
    for j, w in enumerate(larguras, 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = "A2"

r_ = lambda v: v / 100  # centavos -> reais

aba("Resumo e Fluxo de Caixa", [
    ("Indicador", MES_LABEL),
    ("Entradas (realizado)", r_(entradas_mes)),
    ("Saídas (realizado)", r_(saidas_mes)),
    ("Resultado do mês", r_(resultado_mes)),
    (f"Saldo em {FIM_LABEL}", r_(saldo_final)),
    ("", None),
    ("Saldo anterior", r_(saldo_anterior)),
    ("(+) Total de entradas", r_(entradas_mes)),
    ("(-) Total de saídas", r_(saidas_mes)),
    ("Resultado do mês", r_(resultado_mes)),
    ("Saldo final", r_(saldo_final)),
], [28, 18])

aba("Categorias", 
    [("Categoria", "Entradas", "Saídas")] +
    [(n, r_(v), None) for n, v in sorted(rec_vals.items(), key=lambda x: -x[1])] +
    [(n, None, r_(v)) for n, v in sorted(desp_cat.items(), key=lambda x: x[1])] +
    [("Totais", r_(entradas_mes), r_(saidas_mes)),
     ("Resultado", None, r_(resultado_mes))],
    [45, 16, 16])

def aba_dre(nome, receita, grupos, resultado):
    linhas = [("Conta", MES_LABEL, "% receita"), ("Receita total", r_(receita), 1.0)]
    for g, v in grupos:
        linhas.append((f"(-) {GRUPO_LABEL[g]}", r_(v), abs(v) / receita))
    linhas.append((f"Resultado do mês{' (competência)' if 'compet' in nome.lower() else ''}", r_(resultado), resultado / receita))
    aba(nome, linhas, [40, 18, 12])
    ws = wb[nome]
    for row in ws.iter_rows(min_row=3, min_col=3):
        for c in row:
            if isinstance(c.value, float):
                c.number_format = "0.0%"

aba_dre("DRE Caixa", entradas_mes,
        [(g, grupo_val[g]) for g in GRUPOS_ORDEM if grupo_val.get(g)],
        resultado_mes)
aba_dre("DRE Competência", comp_te,
        [(g, sum(val for cat, val in comp_desp.items() if cat_grupo_id.get(cat) == g or cat == g))
         for g in GRUPOS_COMP if any(cat_grupo_id.get(cat) == g or cat == g for cat in comp_desp)],
        comp_te + comp_ts)

aba("Receitas em Aberto",
    [("Cliente", "Valor em aberto")] +
    [(n, r_(v)) for n, v in sorted(por_cliente.items(), key=lambda x: -x[1])] +
    [("Total", r_(total_aberto))],
    [50, 18])

aba("Saldos",
    [("Conta", f"Saldo em {FIM_LABEL}")] +
    [(n, r_(v)) for n, v in saldos_conta] +
    [("Total", r_(sum(v for _, v in saldos_conta)))],
    [35, 18])

aba("Projeção 12 meses",
    [("Mês", "Saldo projetado")] +
    [(m, r_(v)) for m, v in proj],
    [12, 18])

wb.save(ARQ_XLSX)
print(f"OK: {ARQ}")
print(f"OK: {ARQ_XLSX}")
