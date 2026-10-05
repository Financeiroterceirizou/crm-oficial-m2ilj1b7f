#!/usr/bin/env python3
# Relatórios MINAS 1 (JF-T) — 7 unidades por CENTRO DE CUSTO — v1.1, 2026-10-05
#   v1.1 (feedback Vinícius): comparativo ordena RECEITAS primeiro, DESPESAS depois (PDF e Excel).
#   Formato Campo Belo/Uberlândia v1.4 (aprovado): relatórios 1+2 em REGIME DE COMPETÊNCIA
#   (janela 1 ano — há recorrências 2027 com dt_competence retroativa), demais no caixa.
#   Filtro: SÓ categoria 99.01 (regra Vinícius 05/10 — descrição não filtra).
#   Unidades (CC): BELO HORIZONTE 2 170226 · BELO HORIZONTE 3 170224 · CONSELHEIRO LAFA 170225 ·
#   JUIZ DE FORA 1 170223 · MANHUACU 170222 · MURIAE 170221 · UBA 170220.
#   Global = conjunto dos 7 CCs. Lançamentos sem CC (transferências 486 + 24 avulsos) ficam fora.
# Uso: python3 relatorio_minas1.py [YYYY-MM-DD] [unidade|global]  (default: hoje, todas)
import json, os, sys, urllib.request, time
from collections import defaultdict
from datetime import date, timedelta

from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.units import cm
from reportlab.lib import colors
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.platypus import (SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle,
                                PageBreak, Image as RLImage, KeepTogether)

BASE = "https://api-v1.controlle.com"
_dir = os.path.dirname(os.path.abspath(__file__))
TOKEN = os.environ.get("CONTROLLE_TOKEN_MINAS1") or open(os.path.join(_dir, ".controlle_token")).read().strip()
UA = {"Authorization": f"Bearer {TOKEN}", "User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) Chrome/126.0"}
LARANJA = colors.HexColor("#ff501c")
LARANJA_CLARO = colors.HexColor("#ffe3d6")
PRETO = colors.HexColor("#1a1a1a")
CINZA = colors.HexColor("#f5f5f5")
VERDE = colors.HexColor("#1a7f37")
VERMELHO = colors.HexColor("#c0392b")
LOGO = os.path.join(_dir, "logo-terceirizou.png")
SIMBOLO = os.path.join(_dir, "simbolo-terceirizou.png")

CLIENTE = "MINAS 1 - JF(T), UBÁ(MW), CONSELHEIRO(LIMA), MURIAÉ(MARINO), BH2(RCL), MANHUAÇU(LS) e BH3 (MS)"
UNIDADES = {
    "bh2":  {"nome": "BELO HORIZONTE 2", "cc": 170226, "contas": ["BELO HORIZONTE 2 CORA", "BELO HORIZONTE 2 CAIXINHA", "BELO HORIZONTE 2 TRANSPOC"]},
    "bh3":  {"nome": "BELO HORIZONTE 3", "cc": 170224, "contas": ["BELO HORIZONTE 3 CORA", "BELO HORIZONTE 3 CAIXINHA", "BELO HORIZONTE 3 TRANSPOC"]},
    "cons": {"nome": "CONSELHEIRO LAFAIETE", "cc": 170225, "contas": ["CONSELHEIRO LAFAIETE CORA", "CONSELHEIRO LAFAIETE CAIX", "CONSELHEIRO LAFAIETE TRAN"]},
    "jf1":  {"nome": "JUIZ DE FORA 1", "cc": 170223, "contas": ["JUIZ DE FORA 1 CORA", "JUIZ DE FORA 1 CAIXINHA", "JUIZ DE FORA 1 TRANSPOCRE"]},
    "man":  {"nome": "MANHUAÇU", "cc": 170222, "contas": ["MANHUACU CORA", "MANHUACU CAIXINHA", "MANHUACU TRANSPOCRED"]},
    "mur":  {"nome": "MURIAÉ", "cc": 170221, "contas": ["MURIAE CORA", "MURIAE CAIXINHA", "MURIAE TRANSPOCRED"]},
    "uba":  {"nome": "UBÁ", "cc": 170220, "contas": ["UBA CORA", "UBA CAIXINHA", "UBA TRANSPOCRED"]},
}
CCS_GLOBAIS = [u["cc"] for u in UNIDADES.values()]
CAT_VISTORIA = ("01.01", "01.02", "01.03", "01.04", "01.05")

MES_PT = {1:"janeiro",2:"fevereiro",3:"março",4:"abril",5:"maio",6:"junho",7:"julho",8:"agosto",9:"setembro",10:"outubro",11:"novembro",12:"dezembro"}
MES_AB = {1:"jan",2:"fev",3:"mar",4:"abr",5:"mai",6:"jun",7:"jul",8:"ago",9:"set",10:"out",11:"nov",12:"dez"}

def req(url, tries=4):
    for i in range(tries):
        try:
            r = urllib.request.Request(url)
            for k, v in UA.items():
                r.add_header(k, v)
            with urllib.request.urlopen(r, timeout=120) as resp:
                return json.loads(resp.read().decode("utf-8", "replace"))
        except urllib.error.HTTPError as e:
            if e.code in (500, 502, 503) and i < tries - 1:
                time.sleep(3 * (i + 1))
                continue
            raise

def tx_list(start, end, **filtros):
    out, page = [], 1
    while True:
        url = (f"{BASE}/transaction/v1/transactions/list?start_date={start}&end_date={end}"
               f"&page={page}&orderBy=date&orderByCardinality=ASC")
        for k, v in filtros.items():
            url += f"&{k}={v}"
        tl = req(url).get("results", {}).get("transactionsList", [])
        out.extend(tl)
        if len(tl) < 100:
            return out
        page += 1

def brl(cents):
    v = cents / 100
    s = f"{abs(v):,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")
    return ("-" if v < 0 else "") + "R$ " + s

def brl_int(cents):
    v = int(round(cents / 100.0))
    s = f"{abs(v):,}".replace(",", ".")
    return ("-" if v < 0 else "") + s

def add_months(d, n):
    m = d.month - 1 + n
    y = d.year + m // 12
    m = m % 12 + 1
    return date(y, m, 1)

def bdate(t):
    return (t.get("dt_billing") or t.get("dt_due") or "")[:10]

def cdate(t):
    return (t.get("dt_competence") or t.get("dt_due") or "")[:10]

def cat_nome(t):
    cats = t.get("apportionments_plan_account") or []
    return (cats[0].get("ds_category") or "?") if cats else "(sem categoria)"

def cc_ids(t):
    return [a.get("cost_center_id") for a in (t.get("apportionments_cost_center") or [])]

def ok_ub(t):
    """Exclui SÓ a categoria 99.01 (regra Vinícius 05/10) — descrição não filtra."""
    for c in (t.get("apportionments_plan_account") or []):
        if (c.get("ds_category") or "").startswith("99.01"):
            return False
    return True

def agrupa_por_categoria(txs, so_negativas=False, so_positivas=False):
    g = defaultdict(lambda: [0, 0])
    for t in txs:
        v = t["value_in_cent"]
        if so_negativas and v >= 0: continue
        if so_positivas and v <= 0: continue
        g[cat_nome(t)][0] += v
        g[cat_nome(t)][1] += 1
    return g

# ===== datas =====
HOJE = date.fromisoformat(sys.argv[1]) if len(sys.argv) > 1 else date.today()
QUAL = sys.argv[2] if len(sys.argv) > 2 else None
MES_ANT = add_months(HOJE.replace(day=1), -1)
FIM_MES_ANT = MES_ANT.replace(day=28) + timedelta(days=4)
FIM_MES_ANT = FIM_MES_ANT - timedelta(days=FIM_MES_ANT.day)
MES_COR = HOJE.replace(day=1)
FIM_MES_COR = MES_COR.replace(day=28) + timedelta(days=4)
FIM_MES_COR = FIM_MES_COR - timedelta(days=FIM_MES_COR.day)
ANO = MES_ANT.year

HOJE_LABEL = HOJE.strftime("%d/%m/%Y")
FIM_MES_ANT_LABEL = FIM_MES_ANT.strftime("%d/%m/%Y")

# ===== dados (JANELA DE 1 ANO — trap das recorrências 2027 com competência 2026) =====
_tx = tx_list(f"{ANO}-01-01", f"{ANO}-12-31")
# dedup por (id, valor, data) — NUNCA por id sozinho: a licença tem lançamentos DIFERENTES
# com o MESMO id_transactions (ex.: 90601741 = salário -450,00 e -2.055,50 em BH2 set/26)
_vistos = set()
_dedup = []
for t in _tx:
    chave = (t.get("id_transactions"), t.get("value_in_cent"), (t.get("date") or "")[:10])
    if chave in _vistos: continue
    _vistos.add(chave)
    _dedup.append(t)
_tx = _dedup

def meses_6():
    out = []
    for i in range(5, -1, -1):
        ini_m = add_months(MES_ANT, -i)
        fim_m = add_months(ini_m, 1) - timedelta(days=1)
        out.append((ini_m.isoformat(), fim_m.isoformat(), f"{MES_AB[ini_m.month]}/{str(ini_m.year)[2:]}"))
    return out
MESES6 = meses_6()

# ===== relatórios por unidade =====
def dados_unidade(cc_list):
    d = {}
    # 1. rec/desp mês anterior — COMPETÊNCIA (inclui não pagos)
    d["ant_tx_comp"] = [t for t in _tx if cc_ids(t) and any(c in cc_list for c in cc_ids(t))
                        and ok_ub(t) and MES_ANT.isoformat() <= cdate(t) <= FIM_MES_ANT.isoformat()]
    d["ant_rec"] = agrupa_por_categoria(d["ant_tx_comp"], so_positivas=True)
    d["ant_desp"] = agrupa_por_categoria(d["ant_tx_comp"], so_negativas=True)
    d["ant_entradas"] = sum(v for v, _ in d["ant_rec"].values())
    d["ant_saidas"] = sum(v for v, _ in d["ant_desp"].values())
    d["ant_resultado"] = d["ant_entradas"] + d["ant_saidas"]
    # 2. comparativo 6m — COMPETÊNCIA
    matriz = defaultdict(lambda: defaultdict(int))
    for t in _tx:
        if not ok_ub(t) or not cc_ids(t) or not any(c in cc_list for c in cc_ids(t)): continue
        mes = cdate(t)[:7]
        if mes < MESES6[0][0][:7] or mes > MESES6[-1][1][:7]: continue
        aps = t.get("apportionments_plan_account") or []
        if len(aps) <= 1:
            matriz[cat_nome(t)][mes] += t["value_in_cent"]
        else:
            for c in aps:
                matriz[c.get("ds_category") or "?"][mes] += c.get("value") or 0
    d["matriz_6"] = matriz
    # 3. previsão de despesas mês corrente — CAIXA (pago+pendente)
    prev = [t for t in _tx if cc_ids(t) and any(c in cc_list for c in cc_ids(t)) and ok_ub(t)
            and MES_COR.isoformat() <= bdate(t) <= FIM_MES_COR.isoformat()]
    d["prev_desp"] = agrupa_por_categoria(prev, so_negativas=True)
    d["prev_desp_total"] = sum(v for v, _ in d["prev_desp"].values())
    d["prev_tx"] = prev
    # 5. despesas em aberto até fim mês anterior
    ab = [t for t in _tx if cc_ids(t) and any(c in cc_list for c in cc_ids(t)) and ok_ub(t)
          and t["activity_type"] == 0 and t.get("situation") == 0 and bdate(t) <= FIM_MES_ANT.isoformat()]
    d["desp_aberto"] = ab
    g = agrupa_por_categoria(ab)
    d["g_aberto"] = g
    d["total_aberto"] = sum(v for v, _ in g.values())
    return d

# saldos por conta (acumulado desde 2017 até FIM_MES_ANT)
contas_api = req(f"{BASE}/account/v1/accounts").get("results", [])
contas_nome = {c["ds_account"]: c for c in contas_api}

def saldos_unidade(nomes_contas):
    out = []
    for n in nomes_contas:
        c = contas_nome.get(n)
        if not c: continue
        try:
            b = req(f"{BASE}/transaction/v1/transactions/balances?start_date=2017-01-01&end_date={FIM_MES_ANT.isoformat()}&id_account_main={c['id']}")["results"]
            out.append((c["ds_account"], b["balanceDone"]))
        except Exception:
            out.append((n, 0))
    return out

# faturamento vistoria (01.01–01.05) do mês anterior em competência
def fat_vistoria(d):
    return sum(v for cat, (v, _) in d["ant_rec"].items() if cat[:5] in CAT_VISTORIA)

# dias úteis (feriados nacionais fixos + móveis)
def dias_uteis(ano, mes):
    feriados = {date(ano,1,1), date(ano,4,21), date(ano,5,1), date(ano,9,7), date(ano,10,12),
                date(ano,11,2), date(ano,11,15), date(ano,11,20), date(ano,12,25), date(ano,12,31)}
    pascoa = {2026: date(2026,4,5), 2027: date(2027,3,28), 2028: date(2028,4,16)}
    if ano in pascoa:
        p = pascoa[ano]
        feriados |= {p - timedelta(days=48), p - timedelta(days=47), p + timedelta(days=60)}
    d, n = date(ano, mes, 1), 0
    while d.month == mes:
        if d.weekday() < 5 and d not in feriados:
            n += 1
        d += timedelta(days=1)
    return n

print("dados prontos")

# ===== geração =====
styles = getSampleStyleSheet()
h1 = ParagraphStyle("h1", parent=styles["Heading1"], fontSize=16, textColor=PRETO, spaceAfter=4)
h2 = ParagraphStyle("h2", parent=styles["Heading2"], fontSize=12, textColor=LARANJA, spaceBefore=14, spaceAfter=3)
h3 = ParagraphStyle("h3", parent=styles["Heading3"], fontSize=10, textColor=PRETO, spaceBefore=10, spaceAfter=3)
sub = ParagraphStyle("sub", parent=styles["Normal"], fontName="Helvetica-Oblique", fontSize=8.5, textColor=colors.HexColor("#777777"), spaceAfter=10)
cell = ParagraphStyle("cell", parent=styles["Normal"], fontSize=8)
cellb = ParagraphStyle("cellb", parent=styles["Normal"], fontSize=8, fontName="Helvetica-Bold")
cellr = ParagraphStyle("cellr", parent=cell, alignment=2)
cellrb = ParagraphStyle("cellrb", parent=cellb, alignment=2)
cellc = ParagraphStyle("cellc", parent=cell, alignment=1)
P = Paragraph

def P_val(cents, style):
    s = ParagraphStyle("v", parent=style, textColor=(VERDE if cents > 0 else VERMELHO if cents < 0 else PRETO))
    return P(brl(cents), s)

def P_val_int(cents, style):
    s = ParagraphStyle("v", parent=style, textColor=(VERDE if cents > 0 else VERMELHO if cents < 0 else PRETO))
    return P(brl_int(cents), s)

def tabela(rows, widths, fs=7.5):
    t = Table(rows, colWidths=widths, repeatRows=1)
    style = [("FONTNAME", (0,0), (-1,-1), "Helvetica"), ("FONTSIZE", (0,0), (-1,-1), fs),
             ("TOPPADDING", (0,0), (-1,-1), 3), ("BOTTOMPADDING", (0,0), (-1,-1), 3),
             ("VALIGN", (0,0), (-1,-1), "MIDDLE"),
             ("BACKGROUND", (0,0), (-1,0), LARANJA), ("TEXTCOLOR", (0,0), (-1,0), colors.white),
             ("FONTNAME", (0,0), (-1,0), "Helvetica-Bold"),
             ("LINEBELOW", (0,-1), (-1,-1), 0.7, LARANJA)]
    for i in range(1, len(rows)):
        if i % 2 == 0:
            style.append(("BACKGROUND", (0,i), (-1,i), CINZA))
    t.setStyle(TableStyle(style))
    return t

def tabela_cat(titulo, grupos, total_label="Total"):
    rows = [[P("<b>Categoria</b>", cell), P("<b>Lançamentos</b>", cellc), P("<b>Valor</b>", cellr)]]
    for nome, (v, n) in sorted(grupos.items(), key=lambda x: x[0]):
        rows.append([P(nome, cell), P(str(n), cellc), P_val(v, cellr)])
    tot = sum(v for v, _ in grupos.values())
    rows.append([P(f"<b>{total_label}</b>", cellrb), P(f"<b>{sum(n for _, n in grupos.values())}</b>", cellc), P_val(tot, cellrb)])
    t = tabela(rows, [11*cm, 2.5*cm, 3*cm])
    t.setStyle(TableStyle([("BACKGROUND", (0,len(rows)-1), (-1,len(rows)-1), LARANJA_CLARO)]))
    return [P(f"<b>{titulo}</b>", h2), t]

def rodape_simbolo(canvas, doc):
    if os.path.exists(SIMBOLO):
        canvas.drawImage(SIMBOLO, 18.2*cm, 1.1*cm, width=1.4*cm, height=1.4*cm*229/560, mask="auto")

def gerar_pdf(chave, nome_uni, d, saldos, arq_pdf):
    doc = SimpleDocTemplate(arq_pdf, pagesize=A4, leftMargin=1.5*cm, rightMargin=1.5*cm,
                            topMargin=1.3*cm, bottomMargin=1.6*cm,
                            title=f"Relatório Gerencial — {nome_uni}")
    E = []
    h1c = ParagraphStyle("h1c", parent=h1, fontSize=13, alignment=1, spaceAfter=1)
    h2c = ParagraphStyle("h2c", parent=h2, fontSize=11, alignment=1, spaceBefore=2, spaceAfter=2)
    subc = ParagraphStyle("subc", parent=sub, fontSize=8, alignment=1, spaceAfter=0)
    if os.path.exists(LOGO):
        img = RLImage(LOGO, width=5.5*cm, height=5.5*cm*561/1600)
        img.hAlign = "CENTER"
        E.append(img)
    E.append(Spacer(1, 8))
    E.append(P(f"<b>{nome_uni}</b>", h1c))
    E.append(P("Relatório Gerencial Mensal", h2c))
    E.append(P(f"Gerado em {HOJE_LABEL} · Fonte: Controlle · Ref.: {MES_PT[MES_ANT.month].capitalize()} de {MES_ANT.year}", subc))
    E.append(Spacer(1, 6))

    # 1. rec/desp mês anterior (competência)
    E += tabela_cat(f"Receitas e Despesas — {MES_PT[MES_ANT.month].capitalize()} de {MES_ANT.year} (regime de competência)", d["ant_rec"], total_label="Total de Receitas")
    rd_rows = [[P("<b>Categoria</b>", cell), P("<b>Lançamentos</b>", cellc), P("<b>Valor</b>", cellr)]]
    for nome, (v, n) in sorted(d["ant_desp"].items()):
        rd_rows.append([P(nome, cell), P(str(n), cellc), P_val(v, cellr)])
    rd_rows.append([P("<b>Total de Despesas</b>", cellrb), P(f"<b>{sum(n for _, n in d['ant_desp'].values())}</b>", cellc), P_val(d["ant_saidas"], cellrb)])
    rd_rows.append([P("<b>Resultado do mês</b>", cellrb), P(f"<b>{len(d['ant_tx_comp'])}</b>", cellc), P_val(d["ant_resultado"], cellrb)])
    t = tabela(rd_rows, [11*cm, 2.5*cm, 3*cm])
    t.setStyle(TableStyle([("BACKGROUND", (0,len(rd_rows)-2), (-1,len(rd_rows)-1), LARANJA_CLARO)]))
    E.append(t)

    # 2. comparativo 6m (competência) — receitas primeiro, despesas depois
    E.append(Spacer(1, 14))
    E.append(P(f"Comparativo dos Últimos 6 Meses por Categoria — regime de competência", h2))
    cat_names = sorted({c for c in d["matriz_6"]}, key=lambda c: sum(d["matriz_6"][c].values()), reverse=True)
    n_col = len(MESES6)
    cm_rows = [[P("<b>Categoria</b>", cell)] + [P(f"<b>{lab}</b>", cellr) for _, _, lab in MESES6] + [P("<b>Média</b>", cellr)]]
    for cat in cat_names:
        row = [P(cat, cell)]
        vals = [d["matriz_6"][cat].get(fim_m_iso[:7], 0) for _, fim_m_iso, _ in MESES6]
        for v in vals:
            row.append(P_val_int(v, cellr) if v else P("—", cellr))
        row.append(P_val_int(round(sum(vals) / len(vals)), cellrb))
        cm_rows.append(row)
    res_row = [P("<b>Resultado do mês</b>", cellb)]
    _res_mensais = []
    for _, fim_m_iso, _ in MESES6:
        tot_mes = sum(v[fim_m_iso[:7]] for v in d["matriz_6"].values())
        _res_mensais.append(tot_mes)
        res_row.append(P_val_int(tot_mes, cellrb))
    res_row.append(P_val_int(round(sum(_res_mensais) / len(_res_mensais)), cellrb))
    cm_rows.append(res_row)
    E.append(tabela(cm_rows, [6.4*cm] + [1.55*cm]*n_col + [1.6*cm], fs=6.5))

    # 3. previsão de despesas mês corrente
    E.append(PageBreak())
    rows3 = [[P("<b>Categoria</b>", cell), P("<b>Lançamentos</b>", cellc), P("<b>Valor</b>", cellr)]]
    for nm, (v, n) in sorted(d["prev_desp"].items()):
        rows3.append([P(nm, cell), P(str(n), cellc), P_val(v, cellr)])
    rows3.append([P("<b>Total de Despesas Previstas</b>", cellrb), P(f"<b>{sum(n for _, n in d['prev_desp'].values())}</b>", cellc), P_val(d["prev_desp_total"], cellrb)])
    t = tabela(rows3, [11*cm, 2.5*cm, 3*cm])
    t.setStyle(TableStyle([("BACKGROUND", (0,len(rows3)-1), (-1,len(rows3)-1), LARANJA_CLARO)]))
    E.append(P(f"Previsão de Despesas — {MES_PT[MES_COR.month].capitalize()} de {MES_COR.year} (pago + pendente)", h2))
    E.append(t)

    # 4. saldo nas contas
    E.append(Spacer(1, 14))
    E.append(P(f"Saldo nas contas em {FIM_MES_ANT_LABEL}", h2))
    sc_rows = [[P("<b>Conta</b>", cell), P(f"<b>Saldo em {FIM_MES_ANT_LABEL}</b>", cellr)]]
    for nome, v in saldos:
        sc_rows.append([P(nome, cell), P_val(v, cellr)])
    sc_rows.append([P("<b>Total</b>", cellrb), P_val(sum(v for _, v in saldos), cellrb)])
    t = tabela(sc_rows, [11*cm, 5*cm])
    t.setStyle(TableStyle([("BACKGROUND", (0,len(sc_rows)-1), (-1,len(sc_rows)-1), LARANJA_CLARO)]))
    E.append(t)

    # 5. despesas em aberto
    E.append(Spacer(1, 14))
    E.append(P(f"Despesas em aberto até {FIM_MES_ANT_LABEL}", h2))
    E += tabela_cat("Resumo por categoria", d["g_aberto"], total_label="Total em aberto")
    if d["desp_aberto"]:
        por_cat = defaultdict(list)
        for t in d["desp_aberto"]:
            por_cat[cat_nome(t)].append(t)
        E.append(P("<b>Demonstrativo dos lançamentos</b>", h2))
        for cat in sorted(por_cat):
            txs_cat = sorted(por_cat[cat], key=bdate)
            rows = [[P("<b>Vencimento</b>", cell), P("<b>Descrição</b>", cell), P("<b>Conta</b>", cell), P("<b>Situação</b>", cellc), P("<b>Valor</b>", cellr)]]
            for t in txs_cat:
                dd = bdate(t)
                rows.append([P(f"{dd[8:10]}/{dd[5:7]}/{dd[:4]}", cell), P((t.get("ds_transaction") or "")[:70], cell),
                             P(t.get("ds_account_main") or "", cell), P("Aberto", cellc), P_val(t["value_in_cent"], cellr)])
            rows.append([P("<b>Total</b>", cellrb), P(f"<b>{cat}</b>", cellrb), P("", cell), P(f"<b>{len(txs_cat)}</b>", cellc), P_val(sum(t["value_in_cent"] for t in txs_cat), cellrb)])
            tt = tabela(rows, [2.2*cm, 7.3*cm, 3.2*cm, 1.8*cm, 3*cm])
            tt.setStyle(TableStyle([("BACKGROUND", (0,len(rows)-1), (-1,len(rows)-1), LARANJA_CLARO)]))
            E.append(P(f"<b>{cat}</b>", h3))
            E.append(tt)

    # 6. resumo previsão de resultado
    E.append(Spacer(1, 14))
    fat_ant = fat_vistoria(d)
    du_ant = dias_uteis(MES_ANT.year, MES_ANT.month)
    du_cor = dias_uteis(MES_COR.year, MES_COR.month)
    fat_prev = round(fat_ant / du_ant * du_cor) if du_ant else 0
    res_prev = fat_prev + d["prev_desp_total"]
    resumo_rows = [[P("<b>Indicador</b>", cell), P("<b>Valor</b>", cellr)],
                   [P(f"Faturamento {MES_PT[MES_ANT.month].capitalize()}/26 (receitas de vistoria 01.01–01.05, competência)", cell), P_val(fat_ant, cellr)],
                   [P(f"Dias úteis {MES_AB[MES_ANT.month]}/26 → {MES_AB[MES_COR.month]}/26", cell), P(f"{du_ant} → {du_cor}", cellr)],
                   [P(f"Faturamento previsto {MES_PT[MES_COR.month].capitalize()}/26", cell), P_val(fat_prev, cellr)],
                   [P(f"Despesas previstas {MES_PT[MES_COR.month].capitalize()}/26", cell), P_val(d["prev_desp_total"], cellr)],
                   [P("<b>Resultado previsto do mês</b>", cellrb), P_val(res_prev, cellrb)],
                   [P(f"Despesas em aberto até {FIM_MES_ANT_LABEL}", cell), P_val(d["total_aberto"], cellr)],
                   [P(f"Saldo nas contas em {FIM_MES_ANT_LABEL}", cell), P_val(sum(v for _, v in saldos), cellr)],
                   [P("<b>Previsão de saldo final</b> (resultado − desp. em aberto + saldo)", cellrb), P_val(res_prev - d["total_aberto"] + sum(v for _, v in saldos), cellrb)]]
    t = tabela(resumo_rows, [11.5*cm, 4.5*cm])
    t.setStyle(TableStyle([("BACKGROUND", (0,5), (-1,5), LARANJA_CLARO), ("BACKGROUND", (0,8), (-1,8), LARANJA_CLARO)]))
    E.append(P(f"Previsão de Resultado — {MES_PT[MES_COR.month].capitalize()} de {MES_COR.year}", h2))
    E.append(t)
    E.append(P("Base do cálculo: faturamento previsto = média por dia útil do mês anterior (só receitas de vistoria 01.01–01.05) × dias úteis do mês corrente (feriados nacionais descontados); despesas previstas = Previsão de Despesas do mês corrente.", sub))

    E.append(Spacer(1, 10))
    E.append(P("Gerado automaticamente pela Terceirizou · dados do Controlle", sub))
    doc.build(E, onFirstPage=rodape_simbolo, onLaterPages=rodape_simbolo)
    print(f"OK: {arq_pdf}")

# ===== Excel =====
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

FILL_H = PatternFill("solid", fgColor="FF501C")
FILL_T = PatternFill("solid", fgColor="FFE3D6")
FH = Font(bold=True, color="FFFFFF")
FB = Font(bold=True)
VERDE_XL = "1A7F37"
VERMELHO_XL = "C0392B"
TOT_LABELS = ("Total", "Totais", "Resultado", "Previsão", "Total de Receitas", "Total de Despesas",
              "Total de Despesas Previstas", "Total previsto", "Total em aberto", "Faturamento previsto")

def aba(wb, nome, linhas, larguras, titulo=None):
    ws = wb.create_sheet(nome[:31])
    if titulo:
        ws.append((titulo,))
        ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=max(1, len(larguras)))
        c = ws.cell(row=1, column=1)
        c.font = Font(bold=True, size=12, color="FF501C")
        c.alignment = Alignment(horizontal="center", vertical="center")
        ws.row_dimensions[1].height = 22
        desloc = 1
    else:
        desloc = 0
    for r in linhas:
        ws.append(list(r))
    hdr_row = desloc + 1
    headers = [str(ws.cell(row=hdr_row, column=j).value or "") for j in range(1, ws.max_column + 1)]
    for c in ws[hdr_row]:
        c.fill = FILL_H
        c.font = FH
        c.alignment = Alignment(horizontal="center", vertical="center")
    for row in ws.iter_rows(min_row=hdr_row + 1):
        label = str(row[0].value or "")
        is_total = label.startswith(TOT_LABELS)
        if is_total:
            for c in row:
                c.fill = FILL_T
                c.alignment = Alignment(horizontal="center", vertical="center")
            row[0].font = FB
        for j, c in enumerate(row[1:], 2):
            if isinstance(c.value, (int, float)):
                hdr = headers[j-1] if j-1 < len(headers) else ""
                if "Lançamentos" in hdr:
                    c.number_format = "0"
                    c.font = Font(bold=is_total)
                else:
                    c.number_format = '"R$" #,##0.00'
                    cor = VERDE_XL if c.value > 0 else (VERMELHO_XL if c.value < 0 else "1A1A1A")
                    c.font = Font(bold=is_total, color=cor)
    for j, w in enumerate(larguras, 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = ws.cell(row=hdr_row + 1, column=1)

r_ = lambda v: v / 100

def gerar_excel(chave, nome_uni, d, saldos, arq_xlsx):
    wb = Workbook()
    wb.remove(wb.active)
    fat_ant = fat_vistoria(d)
    du_ant = dias_uteis(MES_ANT.year, MES_ANT.month)
    du_cor = dias_uteis(MES_COR.year, MES_COR.month)
    fat_prev = round(fat_ant / du_ant * du_cor) if du_ant else 0
    res_prev = fat_prev + d["prev_desp_total"]
    # comparativo 6m — receitas primeiro, despesas depois
    cat_names = sorted({c for c in d["matriz_6"]}, key=lambda c: sum(d["matriz_6"][c].values()), reverse=True)
    _res_mensais = [sum(v[fim_m_iso[:7]] for v in d["matriz_6"].values()) for _, fim_m_iso, _ in MESES6]
    aba(wb, "Comparativo 6m",
        [("Categoria",) + tuple(lab for _, _, lab in MESES6) + ("Média",)] +
        [(cat,) + tuple((r_(d["matriz_6"][cat].get(fim_m_iso[:7], 0)) if d["matriz_6"][cat].get(fim_m_iso[:7], 0) else None) for _, fim_m_iso, _ in MESES6)
         + (r_(round(sum(d["matriz_6"][cat].values()) / len(MESES6))),) for cat in cat_names] +
        [("Resultado do mês",) + tuple(int(round(m / 100)) if m else None for m in _res_mensais)
         + (int(round(sum(_res_mensais) / len(_res_mensais) / 100)),)],
        [40] + [13]*len(MESES6) + [13], titulo=f"Comparativo dos Últimos 6 Meses por Categoria (regime de competência)")
    # mês anterior
    aba(wb, "Receitas e Despesas " + MES_AB[MES_ANT.month],
        [("Categoria", "Lançamentos", "Valor")] +
        [(n, n2, r_(v)) for n, (v, n2) in sorted(d["ant_rec"].items())] +
        [("Total de Receitas", sum(n for _, n in d["ant_rec"].values()), r_(d["ant_entradas"]))] +
        [(n, n2, r_(v)) for n, (v, n2) in sorted(d["ant_desp"].items())] +
        [("Total de Despesas", sum(n for _, n in d["ant_desp"].values()), r_(d["ant_saidas"])),
         ("Resultado do mês", len(d["ant_tx_comp"]), r_(d["ant_resultado"]))],
        [45, 14, 16], titulo=f"Receitas e Despesas — {MES_PT[MES_ANT.month].capitalize()} de {MES_ANT.year} (regime de competência)")
    # previsão
    aba(wb, "Previsão " + MES_AB[MES_COR.month],
        [("Categoria", "Lançamentos", "Valor")] +
        [(n, n2, r_(v)) for n, (v, n2) in sorted(d["prev_desp"].items())] +
        [("Total de Despesas Previstas", sum(n for _, n in d["prev_desp"].values()), r_(d["prev_desp_total"]))],
        [45, 14, 16], titulo=f"Previsão de Despesas — {MES_PT[MES_COR.month].capitalize()} de {MES_COR.year} (pago + pendente)")
    # saldos
    aba(wb, "Saldo contas",
        [("Conta", f"Saldo em {FIM_MES_ANT_LABEL}")] +
        [(n, r_(v)) for n, v in saldos] + [("Total", r_(sum(v for _, v in saldos)))],
        [30, 18], titulo=f"Saldo nas contas em {FIM_MES_ANT_LABEL}")
    # despesas em aberto
    aba(wb, "Despesas em aberto",
        [("Categoria", "Lançamentos", "Valor")] +
        [(n, n2, r_(v)) for n, (v, n2) in sorted(d["g_aberto"].items())] +
        [("Total em aberto", len(d["desp_aberto"]), r_(d["total_aberto"]))],
        [45, 14, 16], titulo=f"Despesas em aberto até {FIM_MES_ANT_LABEL}")
    # resumo
    aba(wb, "Previsão " + MES_AB[MES_COR.month] + " res",
        [("Indicador", "Valor"),
         (f"Faturamento {MES_AB[MES_ANT.month]}/26 (vistoria 01.01–01.05, competência)", r_(fat_ant)),
         (f"Dias úteis {MES_AB[MES_ANT.month]} → {MES_AB[MES_COR.month]}", f"{du_ant} → {du_cor}"),
         (f"Faturamento previsto {MES_AB[MES_COR.month]}/26", r_(fat_prev)),
         (f"Despesas previstas {MES_AB[MES_COR.month]}/26", r_(d["prev_desp_total"])),
         ("Resultado previsto do mês", r_(res_prev)),
         (f"Despesas em aberto até {FIM_MES_ANT_LABEL}", r_(d["total_aberto"])),
         (f"Saldo nas contas em {FIM_MES_ANT_LABEL}", r_(sum(v for _, v in saldos))),
         ("Previsão de saldo final", r_(res_prev - d["total_aberto"] + sum(v for _, v in saldos)))],
        [52, 18], titulo=f"Previsão de Resultado — {MES_PT[MES_COR.month].capitalize()} de {MES_COR.year}")
    wb.save(arq_xlsx)
    print(f"OK: {arq_xlsx}")

# ===== execução =====
ALVOS = [QUAL] if QUAL in UNIDADES or QUAL == "global" else list(UNIDADES.keys()) + ["global"]
for chave in ALVOS:
    if chave == "global":
        nome_uni = "MINAS 1 — CONSOLIDADO (7 UNIDADES)"
        cc_list = CCS_GLOBAIS
        nomes_contas = [n for u in UNIDADES.values() for n in u["contas"]]
        arq_pdf = f"artifacts/{HOJE.strftime('%y%m%d')}_Relatorio_Minas1_Global.pdf"
        arq_xlsx = f"artifacts/{HOJE.strftime('%y%m%d')}_Relatorio_Minas1_Global.xlsx"
    else:
        u = UNIDADES[chave]
        nome_uni = u["nome"]
        cc_list = [u["cc"]]
        nomes_contas = u["contas"]
        arq_pdf = f"artifacts/{HOJE.strftime('%y%m%d')}_Relatorio_Minas1_{chave.upper()}.pdf"
        arq_xlsx = f"artifacts/{HOJE.strftime('%y%m%d')}_Relatorio_Minas1_{chave.upper()}.xlsx"
    d = dados_unidade(cc_list)
    saldos = saldos_unidade(nomes_contas)
    gerar_pdf(chave, nome_uni, d, saldos, arq_pdf)
    gerar_excel(chave, nome_uni, d, saldos, arq_xlsx)
print("TUDO OK")
